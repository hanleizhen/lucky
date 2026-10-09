"""Targeted Excel writing for the supplied CNC program-sheet template."""

from __future__ import annotations

import re
from copy import copy
from dataclasses import dataclass
from datetime import date
from pathlib import Path
from typing import Mapping, Sequence

from openpyxl import load_workbook
from openpyxl.drawing.image import Image as ExcelImage
from openpyxl.drawing.spreadsheet_drawing import AnchorMarker, OneCellAnchor
from openpyxl.drawing.xdr import XDRPositiveSize2D
from openpyxl.utils.cell import column_index_from_string, coordinate_from_string
from openpyxl.utils.units import pixels_to_EMU

from .models import ImagePlacement, ProgramRecord


class TemplateError(RuntimeError):
    pass


@dataclass(frozen=True, slots=True)
class TemplateLayout:
    sheet_name: str = "Sheet1"
    # The empty merged field at the top-right of the supplied template is the
    # operator's job/program name.  It is also used as the exported workbook
    # filename prefix when it has been filled in.
    output_name_cell: str = "J1"
    # The supplied template labels G3 as DATE and uses J3 for its value.
    date_cell: str = "J3"
    first_program_row: int = 7
    # Rows 26–31 are merged signature/approval areas, not NC-program rows.
    last_program_row: int = 25
    # Column names are intentionally fixed to the supplied template, not inferred.
    columns: Mapping[str, str] = None  # type: ignore[assignment]

    def __post_init__(self) -> None:
        if self.columns is None:
            object.__setattr__(
                self,
                "columns",
                {
                    "program_name": "B",
                    "tool_type": "C",
                    "tool_number": "I",
                    "diameter": "D",
                    "radius": "E",
                    "machining_data": "F",
                    "depth": "J",
                },
            )

    @property
    def capacity(self) -> int:
        return self.last_program_row - self.first_program_row + 1

    def target_cell(self, field: str, index: int) -> str:
        if index < 0 or index >= self.capacity:
            raise TemplateError(f"模板最多可填写 {self.capacity} 个 NC 程序。")
        return f"{self.columns[field]}{self.first_program_row + index}"

    def record_values(self, record: ProgramRecord) -> dict[str, str]:
        return {
            "program_name": record.program_name,
            "tool_type": record.tool_type,
            "tool_number": record.tool_number,
            "diameter": record.diameter,
            "radius": record.radius,
            "machining_data": record.machining_data,
            "depth": record.depth,
        }


LAYOUT = TemplateLayout()


def verify_template(path: str | Path) -> None:
    """Fail early when a user selects a workbook with a different structure."""

    template = Path(path)
    if not template.is_file():
        raise TemplateError(f"模板不存在：{template}")
    try:
        workbook = load_workbook(template, read_only=False, data_only=False, keep_links=True)
    except Exception as exc:  # openpyxl errors vary by file corruption type
        raise TemplateError(f"无法读取 Excel 模板：{exc}") from exc
    try:
        if LAYOUT.sheet_name not in workbook.sheetnames:
            raise TemplateError(f"模板中未找到工作表“{LAYOUT.sheet_name}”。")
        sheet = workbook[LAYOUT.sheet_name]
        if sheet["B5"].value != "File name" or sheet["B6"].value != "NC name":
            raise TemplateError("所选文件不是当前支持的 CNC程序单 模板（B5/B6 表头不匹配）。")
        if sheet.max_row < LAYOUT.last_program_row:
            raise TemplateError("模板的 NC 程序填写区不完整。")
    finally:
        workbook.close()


def program_sheet_date(value: date | None = None) -> str:
    """Return the date in the exact display convention used by the template."""

    return (value or date.today()).strftime("%d / %m / %Y")


def automatic_cells(records: Sequence[ProgramRecord], generated_on: date | None = None) -> dict[str, str]:
    """Values automatically maintained for a generated program sheet."""

    if len(records) > LAYOUT.capacity:
        raise TemplateError(f"一次最多导入 {LAYOUT.capacity} 个 NC 文件；当前为 {len(records)} 个。")
    # J3 is deliberately refreshed for previews and export; the template file
    # itself is never saved or changed.
    values: dict[str, str] = {LAYOUT.date_cell: program_sheet_date(generated_on)}
    for index, record in enumerate(records):
        for field, value in LAYOUT.record_values(record).items():
            values[LAYOUT.target_cell(field, index)] = value
    return values


def automatic_cells_for_rows(
    rows: Sequence[ProgramRecord | None], generated_on: date | None = None
) -> dict[str, str]:
    """Return automatic values for program rows that may contain a manual gap.

    The preview lets an operator insert a blank row between two NC programs.
    A ``None`` item represents that intentional blank row; it is not an
    unrecognised NC file and must not cause later programs to move back up.
    """

    if len(rows) > LAYOUT.capacity:
        raise TemplateError(f"一次最多可填写 {LAYOUT.capacity} 个 NC 文件；当前为 {len(rows)} 个程序行。")
    values: dict[str, str] = {LAYOUT.date_cell: program_sheet_date(generated_on)}
    for index, record in enumerate(rows):
        if record is None:
            continue
        for field, value in LAYOUT.record_values(record).items():
            values[LAYOUT.target_cell(field, index)] = value
    return values


def cleared_program_cells() -> dict[str, str]:
    """Blank the imported-program area without altering the source template.

    The supplied template contains completed sample entries. A newly generated
    sheet must not carry those unrelated rows after the imported NC list ends.
    """

    return {
        LAYOUT.target_cell(field, index): ""
        for index in range(LAYOUT.capacity)
        for field in LAYOUT.columns
    }


def output_filename(
    records: Sequence[ProgramRecord],
    desktop: Path,
    document_name: str = "",
    generated_on: date | None = None,
) -> Path:
    """Return a unique output name without ever reusing an existing workbook.

    A value entered in the template's J1 name field takes precedence and is
    followed by the generation date.  Leaving that field blank retains the
    original convenient NC-program based filename behaviour.
    """

    requested_name = document_name.strip()
    # Operators sometimes type the extension out of habit.  The application
    # always appends the one real .xlsx extension, so remove only that suffix.
    if requested_name.lower().endswith(".xlsx"):
        requested_name = requested_name[:-5].rstrip()
    requested_name = re.sub(r'[<>:"/\\\\|?*]', "_", requested_name).strip(". ")
    today = generated_on or date.today()
    if requested_name:
        base_name = f"{requested_name}_{today.isoformat()}"
    else:
        stems = [re.sub(r'[<>:"/\\\\|?*]', "_", Path(record.program_name).stem) for record in records]
        if stems and len(stems) <= 4:
            base_name = "CNC程序单_" + "_".join(stems)
        else:
            base_name = f"CNC程序单_{today.isoformat()}"
    target = desktop / f"{base_name}.xlsx"
    number = 2
    while target.exists():
        target = desktop / f"{base_name}_{number}.xlsx"
        number += 1
    return target


def export_workbook(
    template_path: str | Path,
    target_path: str | Path,
    cell_values: Mapping[str, str],
    images: Sequence[ImagePlacement] = (),
    red_text_cells: Sequence[str] = (),
) -> None:
    """Create a new workbook. The source template is opened read-only in spirit and never saved."""

    verify_template(template_path)
    target = Path(target_path)
    if target.exists():
        raise TemplateError(f"为保护已有文件，拒绝覆盖：{target}")
    target.parent.mkdir(parents=True, exist_ok=True)
    workbook = load_workbook(template_path, read_only=False, data_only=False, keep_links=True)
    try:
        sheet = workbook[LAYOUT.sheet_name]
        for cell, value in cell_values.items():
            # Preview provides A1 coordinates only. Keep this guard as an extra safety net.
            if not re.fullmatch(r"[A-Z]{1,3}[1-9][0-9]*", cell):
                raise TemplateError(f"无效的单元格地址：{cell}")
            sheet[cell] = value
        for cell in red_text_cells:
            if not re.fullmatch(r"[A-Z]{1,3}[1-9][0-9]*", cell):
                raise TemplateError(f"无效的单元格地址：{cell}")
            # Keep every template font property (family, size, borders and
            # alignment) and change only the instruction text to red.
            font = copy(sheet[cell].font)
            font.color = "FFFF0000"
            sheet[cell].font = font
        for placement in images:
            if not re.fullmatch(r"[A-Z]{1,3}[1-9][0-9]*", placement.anchor):
                raise TemplateError(f"图片的单元格地址无效：{placement.anchor}")
            if not placement.source_path.is_file():
                raise TemplateError(f"找不到已插入的图片：{placement.source_path}")
            if placement.width <= 0 or placement.height <= 0:
                raise TemplateError(f"图片尺寸无效：{placement.source_path.name}")
            if placement.offset_x < 0 or placement.offset_y < 0:
                raise TemplateError(f"图片位置无效：{placement.source_path.name}")
            try:
                image = ExcelImage(str(placement.source_path))
            except Exception as exc:
                raise TemplateError(f"无法读取图片“{placement.source_path.name}”：{exc}") from exc
            column_letters, row = coordinate_from_string(placement.anchor)
            image.anchor = OneCellAnchor(
                _from=AnchorMarker(
                    col=column_index_from_string(column_letters) - 1,
                    row=row - 1,
                    colOff=pixels_to_EMU(placement.offset_x),
                    rowOff=pixels_to_EMU(placement.offset_y),
                ),
                ext=XDRPositiveSize2D(
                    cx=pixels_to_EMU(placement.width),
                    cy=pixels_to_EMU(placement.height),
                ),
            )
            sheet.add_image(image)
        workbook.save(target)
    except Exception:
        # A failed save does not touch the source template.  The caller always
        # creates a fresh unique output path, so an incomplete target can be
        # safely removed without ever deleting an existing user file.
        if target.exists():
            target.unlink()
        raise
    finally:
        workbook.close()
