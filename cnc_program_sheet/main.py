from __future__ import annotations

import ctypes
import os
import re
import sys
import traceback
import xml.etree.ElementTree as ElementTree
from dataclasses import dataclass, replace
from pathlib import Path
from typing import Any, Callable

from openpyxl import load_workbook
from openpyxl.cell.cell import MergedCell
from PySide6.QtCore import QBuffer, QEvent, QIODevice, QPoint, QRectF, QThread, QTimer, Qt, Signal
from PySide6.QtGui import QAction, QColor, QFont, QIcon, QImage, QKeySequence, QPageLayout, QPageSize, QPainter, QPalette, QPen, QPixmap
from PySide6.QtWidgets import (
    QAbstractItemView,
    QApplication,
    QCheckBox,
    QComboBox,
    QColorDialog,
    QDialog,
    QDialogButtonBox,
    QFileDialog,
    QFormLayout,
    QFrame,
    QHBoxLayout,
    QLabel,
    QLineEdit,
    QListWidget,
    QListWidgetItem,
    QMainWindow,
    QMenu,
    QMessageBox,
    QPlainTextEdit,
    QPushButton,
    QSplitter,
    QStatusBar,
    QSystemTrayIcon,
    QTableWidget,
    QTableWidgetItem,
    QTableWidgetSelectionRange,
    QTabWidget,
    QVBoxLayout,
    QWidget,
    QSpinBox,
)
from PySide6.QtPrintSupport import QPrinter, QPrinterInfo, QPrintPreviewWidget

from .app_paths import (
    cache_pasted_nc_program,
    cache_user_image_bytes,
    copy_user_image,
    default_user_template,
    load_embedded_update_repository,
    load_settings,
    log_exception,
    resource_path,
    save_settings,
)
from .excel_service import (
    LAYOUT,
    TemplateError,
    automatic_cells_for_rows,
    cleared_program_cells,
    export_workbook,
    output_filename,
    verify_template,
)
from .models import ImagePlacement, ProgramRecord, UNRECOGNIZED
from .nc_parser import parse_file_records, parse_records_text
from .nc_transform import TransformMode, decode_windows_default, encode_windows_default, transform_file
from .update_service import (
    UpdateError,
    UpdateInfo,
    check_latest,
    download_verified_installer,
    open_release_page,
    start_installer_after_exit,
)
from .version import APP_ICON_FILE, APP_NAME, APP_USER_MODEL_ID, __version__


RESULT_COLUMNS: list[tuple[str, str]] = [
    ("程序名称", "program_name"),
    ("区分", "tool_type"),
    ("刀具号", "tool_number"),
    ("D", "diameter"),
    ("R", "radius"),
    ("加工DATA", "machining_data"),
    ("Depth", "depth"),
    ("解析状态", "status"),
]
FIELD_NAMES = {field for _, field in RESULT_COLUMNS if field != "status"}
IMAGE_FILE_SUFFIXES = {".png", ".jpg", ".jpeg", ".bmp", ".gif"}
ROTATION_MARKERS = ("❮Y⟲180°❯", "❮X⟲180°❯", "❮Z⟲180°❯")
MARKER_RED = QColor("#FF0000")
DEFAULT_SELECTION_COLOR = "#B7D7FF"
CONVERSION_TAG_BACKGROUND = QColor("#DDEBF7")
CONVERSION_TAG_FOREGROUND = QColor("#1F4E78")
CONVERSION_TAGS = {
    TransformMode.M06: "M06",
    TransformMode.M304: "M304",
    TransformMode.COPPER: "TG",
}
CONVERSION_TITLES = {
    TransformMode.M06: "M06 程序修改",
    TransformMode.M304: "M304 程序修改",
    TransformMode.COPPER: "铜工程序修改",
}
OPTIONAL_TOOLBAR_BUTTONS: tuple[tuple[str, str], ...] = (
    ("save_location", "保存位置"),
    ("reparse", "重新解析"),
    ("undo", "撤销"),
    ("redo", "重做"),
    ("find_replace", "查找替换"),
    ("clear", "清空"),
    ("template", "打开模板"),
    ("print", "打印"),
    ("m06", "M06"),
    ("m304", "M304"),
    ("copper", "铜工"),
    ("updates", "检查更新"),
)

# Keep the visual system in one place. This is intentionally presentation
# only: parsing, conversion, workbook generation and keyboard dispatch must
# not depend on a style rule. Applying it to QApplication (rather than only
# MainWindow) also makes modal settings, conversion and print dialogs feel
# like parts of the same product.
APPLICATION_STYLE_SHEET = """
QWidget {
    color: #1F2937;
    font-family: "Microsoft YaHei UI", "Segoe UI", Arial;
    font-size: 12px;
}
QMainWindow, QDialog, QMessageBox {
    background: #F6F8FB;
}
QLabel {
    color: #243B53;
}
QPushButton {
    min-height: 28px;
    padding: 3px 12px;
    color: #1F2937;
    background: #FFFFFF;
    border: 1px solid #C9D4E2;
    border-radius: 5px;
}
QPushButton:hover {
    background: #F0F6FF;
    border-color: #76A9E0;
}
QPushButton:pressed {
    background: #DCEBFA;
    border-color: #4D8DCC;
}
QPushButton:disabled {
    color: #94A3B8;
    background: #F1F5F9;
    border-color: #DCE3EB;
}
QPushButton#primaryButton {
    color: #FFFFFF;
    font-weight: 600;
    background: #0F6CBD;
    border-color: #0F6CBD;
}
QPushButton#primaryButton:hover {
    background: #0B5EA8;
    border-color: #0B5EA8;
}
QPushButton#primaryButton:pressed {
    background: #084B87;
    border-color: #084B87;
}
QLineEdit, QPlainTextEdit, QComboBox, QSpinBox {
    min-height: 26px;
    color: #1F2937;
    background: #FFFFFF;
    border: 1px solid #BFCBDA;
    border-radius: 5px;
    padding: 3px 7px;
    selection-background-color: #B7D7FF;
    selection-color: #111827;
}
QLineEdit:focus, QPlainTextEdit:focus, QComboBox:focus, QSpinBox:focus {
    border: 1px solid #2684D9;
}
QPlainTextEdit {
    padding: 6px;
}
QComboBox::drop-down {
    width: 23px;
    border: 0;
}
QComboBox QAbstractItemView {
    color: #1F2937;
    background: #FFFFFF;
    border: 1px solid #BFCBDA;
    selection-background-color: #DCEBFA;
    selection-color: #1F2937;
}
QTableWidget, QListWidget {
    color: #1F2937;
    background: #FFFFFF;
    border: 1px solid #D4DCE7;
    alternate-background-color: #F8FAFC;
    gridline-color: #D4DCE7;
}
QTableWidget::item:selected, QListWidget::item:selected {
    color: #111827;
    background: #DCEBFA;
}
QHeaderView::section {
    min-height: 26px;
    padding: 3px 6px;
    color: #334155;
    font-weight: 600;
    background: #F1F5F9;
    border: 0;
    border-right: 1px solid #D4DCE7;
    border-bottom: 1px solid #C7D2DE;
}
QTabWidget::pane {
    background: #FFFFFF;
    border: 1px solid #D4DCE7;
    border-radius: 5px;
    top: -1px;
}
QTabBar::tab {
    min-height: 27px;
    padding: 4px 14px;
    color: #52657A;
    background: #EAF0F6;
    border: 1px solid #D4DCE7;
    border-bottom: 0;
    border-top-left-radius: 5px;
    border-top-right-radius: 5px;
}
QTabBar::tab:selected {
    color: #0F5FA8;
    font-weight: 600;
    background: #FFFFFF;
}
QMenu {
    padding: 5px;
    color: #1F2937;
    background: #FFFFFF;
    border: 1px solid #BFCBDA;
    border-radius: 5px;
}
QMenu::item {
    min-height: 25px;
    padding: 4px 28px 4px 24px;
    border-radius: 3px;
}
QMenu::item:selected {
    color: #0B4F8A;
    background: #E5F1FC;
}
QMenu::item:disabled {
    color: #94A3B8;
}
QMenu::separator {
    height: 1px;
    margin: 5px 8px;
    background: #E2E8F0;
}
QToolTip {
    padding: 5px 7px;
    color: #FFFFFF;
    background: #243B53;
    border: 1px solid #334E68;
    border-radius: 3px;
}
QStatusBar {
    color: #52657A;
    background: #F8FAFC;
    border-top: 1px solid #DCE3EB;
}
#dropArea {
    border: 2px dashed #4B87B9;
    border-radius: 9px;
    background: #EDF6FD;
    color: #245B85;
    font-size: 15px;
}
#conversionDropArea {
    border: 2px dashed #79A9CF;
    border-radius: 8px;
    background: #EDF6FD;
    color: #245B85;
    font-size: 13px;
}
QTableWidget#excelPreview {
    gridline-color: #000000;
    border: 1px solid #000000;
}
QTableWidget#excelPreview::item {
    border-right: 1px solid #000000;
    border-bottom: 1px solid #000000;
}
QTableWidget#excelPreview QHeaderView::section {
    color: #1F2937;
    font-weight: 400;
    background: #FFFFFF;
    border: 1px solid #000000;
}
"""


def apply_application_theme(application: QApplication) -> None:
    """Install the shared, non-functional desktop visual system."""

    application.setStyleSheet(APPLICATION_STYLE_SHEET)


def natural_program_sort_key(name: str) -> tuple[object, ...]:
    """Sort program names as an operator expects: O-2 before O-10."""

    return tuple(int(part) if part.isdigit() else part for part in re.split(r"(\d+)", name.casefold()))


def _set_windows_app_user_model_id() -> None:
    """Associate the process with the stable Windows shortcut/taskbar identity.

    This is intentionally best-effort: the desktop app still runs on systems
    where the Win32 shell API is unavailable, while Windows gets a stable
    identity before Qt creates its application and taskbar window.
    """

    if sys.platform != "win32":
        return
    try:
        ctypes.windll.shell32.SetCurrentProcessExplicitAppUserModelID(APP_USER_MODEL_ID)
    except (AttributeError, OSError):
        return


def clipboard_nc_paths(mime_data: Any) -> list[Path]:
    """Extract copied NC files from Windows Explorer clipboard data."""

    return [
        Path(url.toLocalFile())
        for url in mime_data.urls()
        if url.isLocalFile() and is_supported_nc_file(Path(url.toLocalFile()))
    ]


def is_nc_program_text(text: str) -> bool:
    """Conservatively distinguish copied NC code from ordinary cell text."""

    if not text or not text.strip():
        return False
    lines = [line.strip() for line in text.splitlines() if line.strip()]
    if not lines:
        return False
    header_found = any(re.match(r"^\(?\s*O(?:[-_A-Z]*\d)[A-Z0-9_-]*\b", line, re.IGNORECASE) for line in lines)
    command_count = sum(
        bool(re.search(r"(?<![A-Z])(?:G\d{1,3}|M\d{1,3}|T\d+)(?!\d)", line, re.IGNORECASE))
        for line in lines
    )
    coordinate_found = any(
        re.search(r"(?<![A-Z])[XYZ]\s*[+-]?(?:\d+\.?\d*|\.\d+)", line, re.IGNORECASE) for line in lines
    )
    # A full copied program normally has an O-number plus command lines. A
    # selected block is still valid when it has multiple NC commands or a
    # motion command with an axis value.
    return (header_found and command_count >= 1) or command_count >= 2 or (command_count >= 1 and coordinate_found)


def clipboard_nc_filename(text: str) -> str:
    """Use the pasted program's O-number when one is present."""

    match = re.search(r"(?im)^\s*\(?\s*(O(?:[-_A-Z]*\d)[A-Z0-9_-]*)\b", text)
    if match:
        return f"{match.group(1).upper()}.NC"
    return "Clipboard-NC.NC"


def is_supported_nc_file(path: Path) -> bool:
    """Accept conventional .NC files and controller programs with no suffix.

    Copper-electrode programs commonly use names such as ``O0788`` without an
    extension.  Files with another explicit suffix remain excluded so an
    accidental spreadsheet/text document cannot be imported as NC code.
    """

    return path.is_file() and (path.suffix.lower() == ".nc" or not path.suffix)


def apply_tool_number_display_policy(records: list[ProgramRecord]) -> list[ProgramRecord]:
    """Hide T numbers for ordinary .NC files, retain them for copper programs.

    The supplied program-sheet convention uses the T column only for
    controller programs that have no filename suffix.  The parser still reads
    the tool number in both cases; this presentation policy is applied only
    after parsing, so it never weakens NC recognition.
    """

    for record in records:
        if record.source_path.suffix.casefold() == ".nc":
            record.tool_number = ""
    return records


class Worker(QThread):
    succeeded = Signal(object)
    failed = Signal(str)

    def __init__(self, task: Callable[[], Any], parent: QWidget | None = None) -> None:
        super().__init__(parent)
        self._task = task

    def run(self) -> None:
        try:
            self.succeeded.emit(self._task())
        except Exception as exc:
            self.failed.emit(str(exc))


@dataclass(slots=True)
class ConversionSource:
    """One NC program staged in a conversion page before it is overwritten."""

    source_path: Path
    text: str


@dataclass(slots=True)
class ConversionResult:
    """One successfully written shortcut output ready for program-sheet import.

    M06 deliberately removes the controller's ``T...M6`` boundaries.  Its
    optional snapshot therefore retains the operator-reviewed text from
    immediately before the exact legacy conversion, allowing a multi-tool
    program to keep one row per original tool section in the program sheet.
    The NC file written to disk is still the unmodified legacy result.
    """

    source_path: Path
    target_path: Path
    parser_snapshot: str | None = None


class NcConversionDropArea(QLabel):
    """Focusable drag/drop target shared by conversion and main import pages."""

    paths_dropped = Signal(object)
    paste_requested = Signal()

    def __init__(
        self,
        accepts_path: Callable[[Path], bool],
        parent: QWidget | None = None,
        *,
        object_name: str = "conversionDropArea",
        instruction: str | None = None,
        minimum_height: int = 74,
    ) -> None:
        super().__init__(parent)
        self._accepts_path = accepts_path
        self.setObjectName(object_name)
        self.setAcceptDrops(True)
        self.setFocusPolicy(Qt.FocusPolicy.StrongFocus)
        self.setAlignment(Qt.AlignmentFlag.AlignCenter)
        self.setMinimumHeight(minimum_height)
        self.setText(instruction or "将 NC 程序拖到这里，或点击此处后按 Ctrl+V 导入\n支持复制的 NC 文件和 NC 程序文本")
        self.setContextMenuPolicy(Qt.ContextMenuPolicy.CustomContextMenu)
        self.customContextMenuRequested.connect(self._show_import_menu)

    def dragEnterEvent(self, event: Any) -> None:
        urls = event.mimeData().urls() if event.mimeData().hasUrls() else []
        if any(self._accepts_path(Path(url.toLocalFile())) for url in urls if url.isLocalFile()):
            event.acceptProposedAction()

    def dragMoveEvent(self, event: Any) -> None:
        urls = event.mimeData().urls() if event.mimeData().hasUrls() else []
        if any(self._accepts_path(Path(url.toLocalFile())) for url in urls if url.isLocalFile()):
            event.acceptProposedAction()

    def dropEvent(self, event: Any) -> None:
        paths = [
            Path(url.toLocalFile())
            for url in event.mimeData().urls()
            if url.isLocalFile() and self._accepts_path(Path(url.toLocalFile()))
        ]
        self.paths_dropped.emit(paths)
        event.acceptProposedAction()

    def mousePressEvent(self, event: Any) -> None:
        # QLabel does not automatically take keyboard focus on a click.  The
        # explicit focus change makes the on-screen "click then Ctrl+V"
        # instruction reliable.
        self.setFocus()
        super().mousePressEvent(event)

    def _show_import_menu(self, position: Any) -> None:
        menu = QMenu(self)
        paste = menu.addAction("粘贴 NC 文件 / 程序（与 Ctrl+V 相同）")
        paste.triggered.connect(self.paste_requested)
        menu.exec(self.mapToGlobal(position))

    def keyPressEvent(self, event: Any) -> None:
        if event.matches(QKeySequence.StandardKey.Paste):
            self.paste_requested.emit()
            event.accept()
            return
        super().keyPressEvent(event)


class NcConversionDialog(QDialog):
    """Editable staging page for one of the three existing NC shortcuts.

    The UI deliberately stages data before the irreversible write.  The
    eventual transformation itself is delegated to ``nc_transform`` and
    therefore preserves the historic shortcut rules, apart from the explicit
    M06 Windows-CRLF output normalisation requested for this application.
    """

    def __init__(self, mode: TransformMode, parent: QWidget | None = None) -> None:
        super().__init__(parent)
        self.mode = mode
        self.converted_files: list[ConversionResult] = []
        self._sources: list[ConversionSource] = []
        self._active_index = -1
        self._loading_editor = False
        self.setWindowTitle(CONVERSION_TITLES[mode])
        self.resize(1020, 650)
        self.setAcceptDrops(True)

        root = QVBoxLayout(self)
        title = QLabel(f"{CONVERSION_TITLES[mode]}（原快捷方式兼容模式）")
        title_font = title.font()
        title_font.setBold(True)
        title_font.setPointSize(title_font.pointSize() + 2)
        title.setFont(title_font)
        root.addWidget(title)

        notice = QLabel(self._notice_text())
        notice.setWordWrap(True)
        notice.setStyleSheet("color: #5b5b5b;")
        root.addWidget(notice)

        self.drop_area = NcConversionDropArea(self._path_is_supported_by_mode)
        self.drop_area.paths_dropped.connect(self.add_paths)
        self.drop_area.paste_requested.connect(self.import_from_clipboard)
        root.addWidget(self.drop_area)

        splitter = QSplitter(Qt.Orientation.Horizontal)
        files_panel = QWidget()
        files_layout = QVBoxLayout(files_panel)
        files_layout.setContentsMargins(0, 0, 0, 0)
        files_layout.addWidget(QLabel("已导入 NC 文件"))
        self.file_list = QListWidget()
        self.file_list.setMinimumWidth(290)
        self.file_list.currentRowChanged.connect(self._select_source)
        self.file_list.installEventFilter(self)
        files_layout.addWidget(self.file_list, 1)
        choose = QPushButton("选择 NC 文件")
        choose.clicked.connect(self.choose_files)
        files_layout.addWidget(choose)
        splitter.addWidget(files_panel)

        editor_panel = QWidget()
        editor_layout = QVBoxLayout(editor_panel)
        editor_layout.setContentsMargins(0, 0, 0, 0)
        self.editor_title = QLabel("NC 程序内容（可直接手动修改）")
        editor_layout.addWidget(self.editor_title)
        self.editor = QPlainTextEdit()
        self.editor.setLineWrapMode(QPlainTextEdit.LineWrapMode.NoWrap)
        self.editor.setPlaceholderText("从左侧选择一个 NC 文件后，可在这里查看和手动修改程序内容。")
        self.editor.setEnabled(False)
        self.editor.textChanged.connect(self._store_editor_text)
        editor_layout.addWidget(self.editor, 1)
        splitter.addWidget(editor_panel)
        splitter.setSizes([350, 650])
        root.addWidget(splitter, 1)

        self.summary = QLabel("尚未导入 NC 文件。")
        root.addWidget(self.summary)
        buttons = QDialogButtonBox(QDialogButtonBox.StandardButton.Cancel)
        self.apply_button = QPushButton("执行修改并填写程序单")
        self.apply_button.setObjectName("primaryButton")
        self.apply_button.clicked.connect(self.apply_conversion)
        buttons.addButton(self.apply_button, QDialogButtonBox.ButtonRole.AcceptRole)
        buttons.rejected.connect(self.reject)
        root.addWidget(buttons)

    def _notice_text(self) -> str:
        if self.mode is TransformMode.M06:
            return "会按原 M6 快捷方式原地覆盖 .NC 文件：任何含有“ M6 ”字样的整行都会删除，并统一写为 Windows CRLF 换行；不建立备份。"
        if self.mode is TransformMode.M304:
            return "会按原 M304 快捷方式原地覆盖文件：在既定 G43/M8 与 M5 位置插入指令，不建立备份。"
        return "会按原铜工快捷方式生成同名无后缀文件并删除原 .NC；安全行和 M30/M99 规则完全保持不变。"

    def _path_is_supported_by_mode(self, path: Path) -> bool:
        if not is_supported_nc_file(path):
            return False
        # These two source shortcuts themselves skip extensionless files.
        return self.mode is TransformMode.M304 or path.suffix.casefold() == ".nc"

    def choose_files(self) -> None:
        filenames, _ = QFileDialog.getOpenFileNames(
            self,
            "选择 NC 程序",
            str(Path.home()),
            "NC 程序 (*.NC *.nc);;铜工无扩展名程序 (*)",
        )
        self.add_paths([Path(filename) for filename in filenames])

    def dragEnterEvent(self, event: Any) -> None:
        urls = event.mimeData().urls() if event.mimeData().hasUrls() else []
        if any(self._path_is_supported_by_mode(Path(url.toLocalFile())) for url in urls if url.isLocalFile()):
            event.acceptProposedAction()

    def dropEvent(self, event: Any) -> None:
        self.add_paths([Path(url.toLocalFile()) for url in event.mimeData().urls() if url.isLocalFile()])
        event.acceptProposedAction()

    def eventFilter(self, watched: Any, event: Any) -> bool:
        if watched is self.file_list and event.type() == QEvent.Type.KeyPress and event.matches(QKeySequence.StandardKey.Paste):
            self.import_from_clipboard()
            event.accept()
            return True
        return super().eventFilter(watched, event)

    def keyPressEvent(self, event: Any) -> None:
        if event.matches(QKeySequence.StandardKey.Paste):
            self.import_from_clipboard()
            event.accept()
            return
        super().keyPressEvent(event)

    def import_from_clipboard(self) -> None:
        mime_data = QApplication.clipboard().mimeData()
        paths = clipboard_nc_paths(mime_data)
        if paths:
            self.add_paths(paths)
            return
        text = mime_data.text() if mime_data.hasText() else ""
        if not is_nc_program_text(text):
            QMessageBox.information(self, "无法导入", "剪贴板中没有可识别的 NC 文件或 NC 程序文本。")
            return
        try:
            self._save_pasted_source(text)
        except Exception as exc:
            QMessageBox.critical(self, "无法导入 NC 程序", str(exc))

    def _save_pasted_source(self, text: str) -> None:
        """Ask for a real source location when clipboard text has no path.

        Explorer-copied files retain their original folders automatically.
        Plain NC text carries no filesystem path at all, so asking once for a
        location is the only way to honour the shortcut's original-folder
        overwrite behaviour rather than silently changing a hidden cache.
        """

        suggested = Path.home() / clipboard_nc_filename(text)
        filename, _ = QFileDialog.getSaveFileName(
            self,
            "选择粘贴 NC 程序的保存位置",
            str(suggested),
            "NC 程序 (*.NC *.nc);;所有文件 (*)",
        )
        if not filename:
            return
        target = Path(filename)
        if self.mode is not TransformMode.M304 and target.suffix.casefold() != ".nc":
            target = target.with_suffix(".NC")
        target.parent.mkdir(parents=True, exist_ok=True)
        target.write_bytes(encode_windows_default(text))
        self._append_source(target, text)
        self.summary.setText(f"已将剪贴板 NC 文本保存到：{target.parent}；执行后会按原快捷方式覆盖该文件。")

    def add_paths(self, paths: list[Path]) -> None:
        skipped = 0
        for path in sorted(paths, key=lambda item: natural_program_sort_key(item.name)):
            if not self._path_is_supported_by_mode(path):
                skipped += 1
                continue
            try:
                self._append_source(path)
            except Exception as exc:
                QMessageBox.critical(self, "无法读取 NC 程序", f"{path.name}\n\n{exc}")
        if skipped:
            self.summary.setText("部分文件未导入：该快捷方式原本只处理 .NC 文件。")

    def _append_source(self, path: Path, text: str | None = None) -> None:
        source_path = path.resolve()
        if any(source.source_path == source_path for source in self._sources):
            return
        source_text = text if text is not None else decode_windows_default(source_path.read_bytes())
        index = len(self._sources)
        self._sources.append(ConversionSource(source_path, source_text))
        item = QListWidgetItem(f"{CONVERSION_TAGS[self.mode]}    {source_path.name}")
        item.setData(Qt.ItemDataRole.UserRole, index)
        item.setBackground(CONVERSION_TAG_BACKGROUND)
        item.setForeground(CONVERSION_TAG_FOREGROUND)
        item_font = item.font()
        item_font.setBold(True)
        item.setFont(item_font)
        self.file_list.addItem(item)
        self.summary.setText(f"已导入 {len(self._sources)} 个程序；可在右侧检查、修改后执行。")
        if self.file_list.currentRow() < 0:
            self.file_list.setCurrentRow(0)

    def _select_source(self, index: int) -> None:
        self._store_editor_text()
        self._active_index = index
        self._loading_editor = True
        try:
            if 0 <= index < len(self._sources):
                source = self._sources[index]
                self.editor.setEnabled(True)
                self.editor.setPlainText(source.text)
                self.editor_title.setText(f"NC 程序内容（可直接修改）：{source.source_path.name}")
            else:
                self.editor.clear()
                self.editor.setEnabled(False)
                self.editor_title.setText("NC 程序内容（可直接手动修改）")
        finally:
            self._loading_editor = False

    def _store_editor_text(self) -> None:
        if not self._loading_editor and 0 <= self._active_index < len(self._sources):
            self._sources[self._active_index].text = self.editor.toPlainText()

    def apply_conversion(self) -> None:
        self._store_editor_text()
        if not self._sources:
            QMessageBox.information(self, "请先导入", "请拖入、选择或粘贴至少一个 NC 程序。")
            return
        converted: list[ConversionResult] = []
        failures: list[str] = []
        for source in self._sources:
            try:
                target = transform_file(source.source_path, self.mode, source_text=source.text)
                if target.is_file():
                    converted.append(
                        ConversionResult(
                            source.source_path,
                            target.resolve(),
                            source.text if self.mode is TransformMode.M06 else None,
                        )
                    )
                else:
                    failures.append(f"{source.source_path.name}：该文件不符合当前快捷方式的处理规则。")
            except Exception as exc:
                failures.append(f"{source.source_path.name}：{exc}")
        if not converted:
            detail = "\n".join(failures) or "没有符合当前快捷方式规则的 NC 文件。"
            QMessageBox.warning(self, "没有可处理的文件", detail)
            return
        self.converted_files = converted
        if failures:
            QMessageBox.warning(
                self,
                "部分程序修改失败",
                "以下程序未能完成修改；已成功修改的程序会继续自动填写到程序单：\n\n" + "\n".join(failures),
            )
        self.accept()


class NcSourceEditorDialog(QDialog):
    """Manual source editor opened from the main imported-file list."""

    def __init__(self, source_path: Path, parent: QWidget | None = None) -> None:
        super().__init__(parent)
        self.source_path = source_path
        self.setWindowTitle(f"编辑 NC 程序：{source_path.name}")
        self.resize(900, 620)
        root = QVBoxLayout(self)
        root.addWidget(QLabel(f"正在编辑：{source_path}\n保存会直接覆盖该 NC 程序文件。"))
        self.editor = QPlainTextEdit()
        self.editor.setLineWrapMode(QPlainTextEdit.LineWrapMode.NoWrap)
        self.original_text = decode_windows_default(source_path.read_bytes())
        self.did_change = False
        self.editor.setPlainText(self.original_text)
        root.addWidget(self.editor, 1)
        buttons = QDialogButtonBox(QDialogButtonBox.StandardButton.Save | QDialogButtonBox.StandardButton.Cancel)
        buttons.accepted.connect(self.save_source)
        buttons.rejected.connect(self.reject)
        root.addWidget(buttons)

    def save_source(self) -> None:
        self.did_change = self.editor.toPlainText() != self.original_text
        if not self.did_change:
            self.accept()
            return
        try:
            self.source_path.write_bytes(encode_windows_default(self.editor.toPlainText()))
        except Exception as exc:
            QMessageBox.critical(self, "无法保存 NC 程序", str(exc))
            return
        self.accept()


class FloatingImageOverlay(QWidget):
    """An Excel-like floating image object drawn above the table viewport."""

    selected = Signal(str)
    geometry_committed = Signal(str, int, int, int, int)
    delete_requested = Signal(str)

    _HANDLE_SIZE = 8
    _MINIMUM_SIZE = 32

    def __init__(self, preview: "ExcelPreview", placement: ImagePlacement) -> None:
        super().__init__(preview.viewport())
        self._preview = preview
        self._placement = placement
        self._pixmap = QPixmap(str(placement.source_path))
        self._selected = False
        self._drag_mode = ""
        self._start_global = QPoint()
        self._start_geometry = (0, 0, placement.width, placement.height)
        self.setMouseTracking(True)
        self.setFocusPolicy(Qt.FocusPolicy.ClickFocus)
        self.setCursor(Qt.CursorShape.OpenHandCursor)
        self.setToolTip("拖动图片可移动；拖动蓝色控制点可调整大小；按 Delete 或点击“删除图片”可删除。")

    @property
    def placement_id(self) -> str:
        return self._placement.placement_id

    @property
    def placement(self) -> ImagePlacement:
        return self._placement

    def set_placement(self, placement: ImagePlacement) -> None:
        self._placement = placement
        self._pixmap = QPixmap(str(placement.source_path))
        self.update()

    def set_selected(self, selected: bool) -> None:
        if self._selected != selected:
            self._selected = selected
            self.update()

    def set_sheet_geometry(self, x: int, y: int, width: int, height: int) -> None:
        view_x, view_y = self._preview.sheet_to_viewport(x, y)
        self.setGeometry(view_x, view_y, max(self._MINIMUM_SIZE, width), max(self._MINIMUM_SIZE, height))

    def sheet_geometry(self) -> tuple[int, int, int, int]:
        x, y = self._preview.viewport_to_sheet(self.pos().x(), self.pos().y())
        return x, y, self.width(), self.height()

    def _handle_at(self, point: QPoint) -> str:
        if not self._selected:
            return "move"
        edge = self._HANDLE_SIZE
        left = point.x() <= edge
        right = point.x() >= self.width() - edge
        top = point.y() <= edge
        bottom = point.y() >= self.height() - edge
        if top and left:
            return "nw"
        if top and right:
            return "ne"
        if bottom and left:
            return "sw"
        if bottom and right:
            return "se"
        if top:
            return "n"
        if bottom:
            return "s"
        if left:
            return "w"
        if right:
            return "e"
        return "move"

    @staticmethod
    def _cursor_for_mode(mode: str) -> Qt.CursorShape:
        if mode in {"nw", "se"}:
            return Qt.CursorShape.SizeFDiagCursor
        if mode in {"ne", "sw"}:
            return Qt.CursorShape.SizeBDiagCursor
        if mode in {"n", "s"}:
            return Qt.CursorShape.SizeVerCursor
        if mode in {"e", "w"}:
            return Qt.CursorShape.SizeHorCursor
        return Qt.CursorShape.OpenHandCursor

    def mousePressEvent(self, event: Any) -> None:
        if event.button() != Qt.MouseButton.LeftButton:
            super().mousePressEvent(event)
            return
        self._selected = True
        self.selected.emit(self.placement_id)
        self._drag_mode = self._handle_at(event.position().toPoint())
        self._start_global = event.globalPosition().toPoint()
        self._start_geometry = self.sheet_geometry()
        self.setCursor(Qt.CursorShape.ClosedHandCursor if self._drag_mode == "move" else self._cursor_for_mode(self._drag_mode))
        self.grabMouse()
        self.setFocus()
        self.update()
        event.accept()

    def mouseMoveEvent(self, event: Any) -> None:
        if not self._drag_mode:
            self.setCursor(self._cursor_for_mode(self._handle_at(event.position().toPoint())))
            event.accept()
            return
        delta = event.globalPosition().toPoint() - self._start_global
        x, y, width, height = self._start_geometry
        mode = self._drag_mode
        if mode == "move":
            x += delta.x()
            y += delta.y()
        else:
            if "e" in mode:
                width += delta.x()
            if "s" in mode:
                height += delta.y()
            if "w" in mode:
                x += delta.x()
                width -= delta.x()
            if "n" in mode:
                y += delta.y()
                height -= delta.y()
        x, y, width, height = self._preview.constrain_image_geometry(x, y, width, height, self._MINIMUM_SIZE)
        self.set_sheet_geometry(x, y, width, height)
        event.accept()

    def mouseReleaseEvent(self, event: Any) -> None:
        if event.button() == Qt.MouseButton.LeftButton and self._drag_mode:
            self.releaseMouse()
            self._drag_mode = ""
            x, y, width, height = self.sheet_geometry()
            self.geometry_committed.emit(self.placement_id, x, y, width, height)
            self.setCursor(Qt.CursorShape.OpenHandCursor)
            event.accept()
            return
        super().mouseReleaseEvent(event)

    def keyPressEvent(self, event: Any) -> None:
        if event.key() == Qt.Key.Key_Delete:
            self.delete_requested.emit(self.placement_id)
            event.accept()
            return
        if event.matches(QKeySequence.StandardKey.Paste):
            self._preview.clipboard_image_paste_requested.emit()
            event.accept()
            return
        super().keyPressEvent(event)

    def paintEvent(self, event: Any) -> None:
        painter = QPainter(self)
        painter.setRenderHint(QPainter.RenderHint.SmoothPixmapTransform)
        if not self._pixmap.isNull():
            painter.drawPixmap(self.rect(), self._pixmap)
        else:
            painter.fillRect(self.rect(), QColor("#FCE4D6"))
            painter.setPen(QColor("#C00000"))
            painter.drawText(self.rect(), Qt.AlignmentFlag.AlignCenter, "无法读取图片")
        if self._selected:
            painter.setPen(QPen(QColor("#2F75B5"), 2))
            painter.drawRect(self.rect().adjusted(1, 1, -2, -2))
            painter.setBrush(QColor("#FFFFFF"))
            painter.setPen(QPen(QColor("#2F75B5"), 1))
            for point in (
                QPoint(0, 0),
                QPoint(self.width() // 2, 0),
                QPoint(self.width() - 1, 0),
                QPoint(0, self.height() // 2),
                QPoint(self.width() - 1, self.height() // 2),
                QPoint(0, self.height() - 1),
                QPoint(self.width() // 2, self.height() - 1),
                QPoint(self.width() - 1, self.height() - 1),
            ):
                painter.drawRect(
                    point.x() - self._HANDLE_SIZE // 2,
                    point.y() - self._HANDLE_SIZE // 2,
                    self._HANDLE_SIZE,
                    self._HANDLE_SIZE,
                )


class ExcelPreview(QTableWidget):
    """Editable in-app preview of the real sheet values and major formatting."""

    cell_edited = Signal(str, str)
    insert_blank_row_requested = Signal(int)
    insert_marker_row_requested = Signal(int, str)
    remove_blank_row_requested = Signal(int)
    clipboard_image_paste_requested = Signal()
    clipboard_nc_paste_requested = Signal()
    clipboard_table_paste_requested = Signal()
    delete_cells_requested = Signal(int, int, int, int)
    image_geometry_changed = Signal(str, str, int, int, int, int)
    image_delete_requested = Signal(str)
    find_replace_requested = Signal(bool)
    print_area_requested = Signal(int, int, int, int)
    clear_print_area_requested = Signal()
    print_area_selection_finished = Signal(int, int, int, int)

    def __init__(self, parent: QWidget | None = None) -> None:
        super().__init__(parent)
        self._loading = False
        self._theme_colors: list[str] = []
        self._image_overlays: dict[str, FloatingImageOverlay] = {}
        self._selected_image_id: str | None = None
        # One history entry may contain an entire pasted/cleared rectangle.
        # This makes Ctrl+Z feel like Excel rather than undoing one cell at a
        # time after a multi-cell paste.
        self._undo_stack: list[list[tuple[str, str, str]]] = []
        self._redo_stack: list[list[tuple[str, str, str]]] = []
        self._known_values: dict[str, str] = {}
        self._history_applying = False
        # Keep enough pointer state to distinguish a normal click from a
        # drag-selection.  A click on an empty grid cell after either one cell
        # or a range has been selected dismisses the visible focus completely.
        self._active_cell_before_click = False
        self._pointer_press_position = QPoint()
        # QTableWidget only contains an item at the top-left of an Excel
        # merged range.  Keep the mapping for every covered preview cell so
        # actions such as Ctrl+V still have a valid Excel anchor.
        self._merged_cell_roots: dict[tuple[int, int], str] = {}
        # Excel cells with “no fill” are white. Do not inherit a dark Windows
        # palette that would make the template preview black.
        palette = self.palette()
        palette.setColor(QPalette.ColorRole.Base, QColor("#FFFFFF"))
        palette.setColor(QPalette.ColorRole.AlternateBase, QColor("#FFFFFF"))
        palette.setColor(QPalette.ColorRole.Text, QColor("#000000"))
        self.setPalette(palette)
        self.setAlternatingRowColors(False)
        self.setWordWrap(True)
        # Keep selection and editing behaviour close to Excel: a mouse drag
        # selects a rectangle, Shift extends it and Ctrl keeps extra cells.
        self.setSelectionMode(QAbstractItemView.SelectionMode.ExtendedSelection)
        self.setSelectionBehavior(QAbstractItemView.SelectionBehavior.SelectItems)
        self.setHorizontalScrollMode(QAbstractItemView.ScrollMode.ScrollPerPixel)
        self.setVerticalScrollMode(QAbstractItemView.ScrollMode.ScrollPerPixel)
        self.setEditTriggers(
            QTableWidget.EditTrigger.DoubleClicked
            | QTableWidget.EditTrigger.EditKeyPressed
            | QTableWidget.EditTrigger.AnyKeyPressed
        )
        # Some Windows palettes use white text for a selected table item.
        # That makes the contents appear to disappear on a light template, so
        # give the selection an explicit Excel-like blue fill and dark text.
        self.set_selection_color(DEFAULT_SELECTION_COLOR)
        self.setContextMenuPolicy(Qt.ContextMenuPolicy.CustomContextMenu)
        self.customContextMenuRequested.connect(self._show_program_row_menu)
        self.itemChanged.connect(self._on_item_changed)
        self.horizontalScrollBar().valueChanged.connect(lambda _: self._reposition_image_overlays())
        self.verticalScrollBar().valueChanged.connect(lambda _: self._reposition_image_overlays())

    @staticmethod
    def _is_program_table_row(row: int) -> bool:
        return LAYOUT.first_program_row - 1 <= row <= LAYOUT.last_program_row - 1

    def _show_program_row_menu(self, position: Any) -> None:
        index = self.indexAt(position)
        if not index.isValid():
            return
        selection = self._selection_rectangle()
        # Calling ``setCurrentCell`` unconditionally clears a multi-cell
        # selection in Qt.  A right-click inside an existing Excel-like range
        # must retain that whole range so the user can make it the print area.
        if selection is None or not (
            selection[0] <= index.row() <= selection[2]
            and selection[1] <= index.column() <= selection[3]
        ):
            self.setCurrentCell(index.row(), index.column())
            selection = (index.row(), index.column(), index.row(), index.column())
        menu = QMenu(self)
        if selection is not None:
            set_print_area = menu.addAction("设为打印区域（记住）")
            set_print_area.triggered.connect(
                lambda: self.print_area_requested.emit(*selection)
            )
        clear_print_area = menu.addAction("清除打印区域（恢复全表）")
        clear_print_area.triggered.connect(self.clear_print_area_requested)
        menu.addSeparator()
        find_replace = menu.addAction("查找/替换…")
        find_replace.triggered.connect(lambda: self.find_replace_requested.emit(True))
        menu.addSeparator()
        paste = menu.addAction("粘贴（与 Ctrl+V 相同；可先在 Win+V 选择内容）")
        paste.triggered.connect(self.request_clipboard_paste)
        copy_action = menu.addAction("复制")
        copy_action.triggered.connect(self.copy_selected_cells)
        cut_action = menu.addAction("剪切")
        cut_action.triggered.connect(self.cut_selected_cells)
        if not self._is_program_table_row(index.row()):
            menu.exec(self.viewport().mapToGlobal(position))
            return
        menu.addSeparator()
        menu.addSection("在下方插入")
        insert = menu.addAction("1. 空行")
        insert.setEnabled(index.row() < LAYOUT.last_program_row - 1)
        insert.triggered.connect(lambda: self.insert_blank_row_requested.emit(index.row() + 1))
        for sequence, marker in enumerate(ROTATION_MARKERS, start=2):
            action = menu.addAction(f"{sequence}. {marker}（红色提示行）")
            action.setIconText(action.text())
            action.setFont(QFont("Microsoft YaHei UI", 10, QFont.Weight.DemiBold))
            action.setEnabled(index.row() < LAYOUT.last_program_row - 1)
            action.triggered.connect(
                lambda checked=False, marker_text=marker: self.insert_marker_row_requested.emit(index.row() + 1, marker_text)
            )
        menu.addSeparator()
        remove = menu.addAction("取消空行（也可按 Delete）")
        remove.triggered.connect(lambda: self.remove_blank_row_requested.emit(index.row()))
        menu.exec(self.viewport().mapToGlobal(position))

    def _selection_rectangle(self) -> tuple[int, int, int, int] | None:
        """Return the bounding rectangle for the current Excel-like selection."""

        indexes = self.selectedIndexes()
        if not indexes:
            return None
        return (
            min(index.row() for index in indexes),
            min(index.column() for index in indexes),
            max(index.row() for index in indexes),
            max(index.column() for index in indexes),
        )

    def selected_cell_range(self) -> tuple[int, int, int, int] | None:
        """Expose the current Excel-like rectangular selection to the window."""

        return self._selection_rectangle()

    @staticmethod
    def _is_plain_click(event: Any, press_position: QPoint) -> bool:
        """Return whether a left-button event ended as a click, not a drag."""

        return (
            event.button() == Qt.MouseButton.LeftButton
            and not (event.modifiers() & (Qt.KeyboardModifier.ControlModifier | Qt.KeyboardModifier.ShiftModifier))
            and (event.position().toPoint() - press_position).manhattanLength() <= 4
        )

    def _is_empty_grid_cell(self, index: Any) -> bool:
        """Treat empty cells and the unused viewport below the sheet as blank."""

        if not index.isValid():
            return True
        item = self.item(index.row(), index.column())
        return item is None or not item.text().strip()

    def set_selection_color(self, color_name: str) -> str:
        """Apply the user's preferred selection fill and matching outline."""

        color = QColor(color_name)
        if not color.isValid():
            color = QColor(DEFAULT_SELECTION_COLOR)
        fill = color.name()
        active_fill = color.darker(108).name()
        outline = color.darker(175).name()
        self.setStyleSheet(
            f"QTableWidget::item:selected {{ background-color: {fill}; color: #000000; }}"
            f"QTableWidget::item:selected:active {{ background-color: {active_fill}; color: #000000; }}"
            f"QTableWidget::item:focus {{ border: 2px solid {outline}; }}"
        )
        return fill

    def _dismiss_range_selection(self) -> None:
        """Return the sheet to a fully normal state without an active cell."""

        self.clearSelection()
        # clearSelection() removes the blue fill but intentionally keeps the
        # active cell.  That cell still draws Qt's blue focus rectangle, which
        # is not a complete cancellation from the operator's point of view.
        self.selectionModel().clearCurrentIndex()
        self.clear_image_selection()
        self.viewport().update()

    def mousePressEvent(self, event: Any) -> None:
        """Remember whether any cell focus existed before a table click."""

        if event.button() == Qt.MouseButton.LeftButton:
            self._active_cell_before_click = self.currentIndex().isValid()
            self._pointer_press_position = event.position().toPoint()
        else:
            self._active_cell_before_click = False
        super().mousePressEvent(event)

    def mouseReleaseEvent(self, event: Any) -> None:
        """Commit a drag-selected print range or dismiss it from a blank click."""

        super().mouseReleaseEvent(event)
        if event.button() != Qt.MouseButton.LeftButton:
            return
        index = self.indexAt(event.position().toPoint())
        if (
            self._active_cell_before_click
            and self._is_plain_click(event, self._pointer_press_position)
            and self._is_empty_grid_cell(index)
        ):
            self._dismiss_range_selection()
            self._active_cell_before_click = False
            return
        self._active_cell_before_click = False
        selection = self._selection_rectangle()
        if selection is not None:
            self.print_area_selection_finished.emit(*selection)

    def _paste_anchor(self) -> tuple[int, int] | None:
        selection = self._selection_rectangle()
        if selection is not None:
            return selection[0], selection[1]
        coordinate = self.current_coordinate()
        if not coordinate:
            return None
        letters = "".join(character for character in coordinate if character.isalpha())
        digits = "".join(character for character in coordinate if character.isdigit())
        if not letters or not digits:
            return None
        column = 0
        for character in letters:
            column = column * 26 + ord(character) - 65 + 1
        return int(digits) - 1, column - 1

    def paste_text_grid(self, text: str) -> int:
        """Paste ordinary clipboard text as a tab/newline cell rectangle."""

        anchor = self._paste_anchor()
        if anchor is None or not text:
            return 0
        start_row, start_column = anchor
        rows = text.replace("\r\n", "\n").replace("\r", "\n").split("\n")
        if rows and rows[-1] == "":
            rows.pop()
        source = [line.split("\t") for line in rows]
        selection = self._selection_rectangle()
        # Excel fills a selected rectangle when a single copied cell is
        # pasted.  Multi-cell clipboard data begins in the top-left cell.
        if len(source) == 1 and len(source[0]) == 1 and selection is not None:
            target_rows = range(selection[0], selection[2] + 1)
            target_columns = range(selection[1], selection[3] + 1)
            grid = [[source[0][0] for _ in target_columns] for _ in target_rows]
            start_row, start_column = selection[0], selection[1]
        else:
            grid = source
        written = 0
        last_row, last_column = start_row, start_column
        changes: list[tuple[str, str, str]] = []
        self._loading = True
        try:
            for row_offset, line in enumerate(grid):
                for column_offset, value in enumerate(line):
                    row, column = start_row + row_offset, start_column + column_offset
                    if row >= self.rowCount() or column >= self.columnCount():
                        continue
                    item = self.item(row, column)
                    if item is None or not (item.flags() & Qt.ItemFlag.ItemIsEditable):
                        continue
                    coordinate = item.data(Qt.ItemDataRole.UserRole)
                    previous = item.text()
                    item.setText(value)
                    if coordinate:
                        coordinate = str(coordinate)
                        if previous != value:
                            changes.append((coordinate, previous, value))
                        self._known_values[coordinate] = value
                        self.cell_edited.emit(str(coordinate), value)
                    written += 1
                    last_row, last_column = max(last_row, row), max(last_column, column)
        finally:
            self._loading = False
        if written:
            self._remember_edit(changes)
            self.clearSelection()
            self.setCurrentCell(start_row, start_column)
            self.setRangeSelected(QTableWidgetSelectionRange(start_row, start_column, last_row, last_column), True)
        return written

    def _selected_range(self) -> tuple[int, int, int, int] | None:
        selection = self._selection_rectangle()
        if selection is None:
            anchor = self._paste_anchor()
            return (*anchor, *anchor) if anchor else None
        return selection

    def copy_selected_cells(self) -> None:
        selection = self._selected_range()
        if selection is None:
            return
        top, left, bottom, right = selection
        lines = []
        for row in range(top, bottom + 1):
            lines.append("\t".join(self.item(row, column).text() if self.item(row, column) is not None else "" for column in range(left, right + 1)))
        QApplication.clipboard().setText("\n".join(lines))

    def cut_selected_cells(self) -> None:
        self.copy_selected_cells()
        self.clear_selected_cells()

    def clear_selected_cells(self) -> None:
        selection = self._selected_range()
        if selection is None:
            return
        top, left, bottom, right = selection
        changes: list[tuple[str, str, str]] = []
        self._loading = True
        try:
            for row in range(top, bottom + 1):
                for column in range(left, right + 1):
                    item = self.item(row, column)
                    if item is not None and item.flags() & Qt.ItemFlag.ItemIsEditable:
                        previous = item.text()
                        item.setText("")
                        coordinate = item.data(Qt.ItemDataRole.UserRole)
                        if coordinate:
                            coordinate = str(coordinate)
                            if previous:
                                changes.append((coordinate, previous, ""))
                            self._known_values[coordinate] = ""
                            self.cell_edited.emit(str(coordinate), "")
        finally:
            self._loading = False
        self._remember_edit(changes)

    def _remember_edit(self, changes: list[tuple[str, str, str]]) -> None:
        """Store one user gesture for Ctrl+Z/Ctrl+Y, omitting no-op cells."""

        effective = [change for change in changes if change[1] != change[2]]
        if not effective or self._history_applying:
            return
        self._undo_stack.append(effective)
        # Keep history bounded so a very large clipboard paste never grows
        # the application's memory indefinitely.
        if len(self._undo_stack) > 100:
            self._undo_stack.pop(0)
        self._redo_stack.clear()

    def clear_edit_history(self) -> None:
        self._undo_stack.clear()
        self._redo_stack.clear()

    def _sync_known_values(self) -> None:
        self._known_values = {
            str(item.data(Qt.ItemDataRole.UserRole)): item.text()
            for row in range(self.rowCount())
            for column in range(self.columnCount())
            if (item := self.item(row, column)) is not None and item.data(Qt.ItemDataRole.UserRole)
        }

    def _item_for_coordinate(self, coordinate: str) -> QTableWidgetItem | None:
        letters = "".join(character for character in coordinate if character.isalpha())
        digits = "".join(character for character in coordinate if character.isdigit())
        if not letters or not digits:
            return None
        column = 0
        for character in letters:
            column = column * 26 + ord(character) - 64
        return self.item(int(digits) - 1, column - 1)

    def _apply_history(self, changes: list[tuple[str, str, str]], use_previous: bool) -> None:
        self._history_applying = True
        self._loading = True
        try:
            ordered = reversed(changes) if use_previous else changes
            for coordinate, previous, current in ordered:
                item = self._item_for_coordinate(coordinate)
                if item is None:
                    continue
                value = previous if use_previous else current
                item.setText(value)
                self._known_values[coordinate] = value
                self.cell_edited.emit(coordinate, value)
        finally:
            self._loading = False
            self._history_applying = False

    def undo_edit(self) -> bool:
        if not self._undo_stack:
            return False
        changes = self._undo_stack.pop()
        self._apply_history(changes, use_previous=True)
        self._redo_stack.append(changes)
        return True

    def redo_edit(self) -> bool:
        if not self._redo_stack:
            return False
        changes = self._redo_stack.pop()
        self._apply_history(changes, use_previous=False)
        self._undo_stack.append(changes)
        return True

    def _editable_items(self) -> list[tuple[int, int, QTableWidgetItem]]:
        """Return editable cells in normal spreadsheet reading order."""

        return [
            (row, column, item)
            for row in range(self.rowCount())
            for column in range(self.columnCount())
            if (item := self.item(row, column)) is not None and item.flags() & Qt.ItemFlag.ItemIsEditable
        ]

    def find_next(self, text: str, match_case: bool = False) -> str | None:
        """Select the next matching editable cell, wrapping at the end."""

        if not text:
            return None
        items = self._editable_items()
        if not items:
            return None
        needle = text if match_case else text.casefold()
        current = (self.currentRow(), self.currentColumn())
        start = next((index for index, (row, column, _) in enumerate(items) if (row, column) == current), -1)
        for offset in range(1, len(items) + 1):
            row, column, item = items[(start + offset) % len(items)]
            haystack = item.text() if match_case else item.text().casefold()
            if needle in haystack:
                self.clearSelection()
                self.setCurrentCell(row, column)
                self.setRangeSelected(QTableWidgetSelectionRange(row, column, row, column), True)
                self.scrollToItem(item, QAbstractItemView.ScrollHint.PositionAtCenter)
                return str(item.data(Qt.ItemDataRole.UserRole) or "")
        return None

    @staticmethod
    def _replace_text(value: str, find_text: str, replacement: str, match_case: bool) -> tuple[str, int]:
        flags = 0 if match_case else re.IGNORECASE
        return re.subn(re.escape(find_text), replacement, value, flags=flags)

    def replace_current(self, find_text: str, replacement: str, match_case: bool = False) -> bool:
        """Replace matches in the active editable cell as one normal edit."""

        item = self.currentItem()
        if not find_text or item is None or not (item.flags() & Qt.ItemFlag.ItemIsEditable):
            return False
        value, count = self._replace_text(item.text(), find_text, replacement, match_case)
        if not count:
            return False
        item.setText(value)
        return True

    def replace_all(self, find_text: str, replacement: str, match_case: bool = False) -> int:
        """Replace across editable cells and keep the full action undoable."""

        if not find_text:
            return 0
        changes: list[tuple[str, str, str]] = []
        replaced_cells = 0
        self._loading = True
        try:
            for _, _, item in self._editable_items():
                previous = item.text()
                current, count = self._replace_text(previous, find_text, replacement, match_case)
                coordinate = item.data(Qt.ItemDataRole.UserRole)
                if not count or not coordinate:
                    continue
                coordinate = str(coordinate)
                item.setText(current)
                self._known_values[coordinate] = current
                self.cell_edited.emit(coordinate, current)
                changes.append((coordinate, previous, current))
                replaced_cells += 1
        finally:
            self._loading = False
        self._remember_edit(changes)
        return replaced_cells

    def request_clipboard_paste(self) -> bool:
        """Dispatch a right-click paste through the exact Ctrl+V policy."""

        mime_data = QApplication.clipboard().mimeData()
        image_file_in_clipboard = any(
            url.isLocalFile() and Path(url.toLocalFile()).suffix.lower() in IMAGE_FILE_SUFFIXES
            for url in mime_data.urls()
        )
        if mime_data.hasImage() or image_file_in_clipboard:
            self.clipboard_image_paste_requested.emit()
            return True
        if clipboard_nc_paths(mime_data) or (mime_data.hasText() and is_nc_program_text(mime_data.text())):
            self.clipboard_nc_paste_requested.emit()
            return True
        if mime_data.hasText():
            self.clipboard_table_paste_requested.emit()
            return True
        return False

    def keyPressEvent(self, event: Any) -> None:
        if event.key() == Qt.Key.Key_Escape and self.currentIndex().isValid():
            # Match the standard spreadsheet escape route as well as the
            # mouse route above.  It deliberately removes the remaining
            # active-cell outline too, returning the sheet to a normal state.
            self._dismiss_range_selection()
            event.accept()
            return
        if event.modifiers() & Qt.KeyboardModifier.ControlModifier and event.key() == Qt.Key.Key_F:
            self.find_replace_requested.emit(False)
            event.accept()
            return
        if event.modifiers() & Qt.KeyboardModifier.ControlModifier and event.key() == Qt.Key.Key_H:
            self.find_replace_requested.emit(True)
            event.accept()
            return
        if event.matches(QKeySequence.StandardKey.Undo):
            self.undo_edit()
            event.accept()
            return
        if event.matches(QKeySequence.StandardKey.Redo):
            self.redo_edit()
            event.accept()
            return
        if event.matches(QKeySequence.StandardKey.Copy):
            self.copy_selected_cells()
            event.accept()
            return
        if event.matches(QKeySequence.StandardKey.Cut):
            self.cut_selected_cells()
            event.accept()
            return
        if event.matches(QKeySequence.StandardKey.Paste) and self.request_clipboard_paste():
            event.accept()
            return
        if event.key() == Qt.Key.Key_Delete:
            # Preserve the earlier shortcut: Delete on an inserted empty row
            # removes that row. Everywhere else Delete behaves like Excel and
            # clears the selected editable cells.
            selected = self._selected_range()
            if selected is not None:
                top, left, bottom, right = selected
                self.delete_cells_requested.emit(top, left, bottom, right)
            event.accept()
            return
        super().keyPressEvent(event)

    def resizeEvent(self, event: Any) -> None:
        super().resizeEvent(event)
        QTimer.singleShot(0, self._reposition_image_overlays)

    def shift_program_row_values(self, start_table_row: int, direction: int) -> None:
        """Move visible program-row values while retaining every cell's style.

        This is a logical row move inside the template's fixed NC area, not an
        Excel row insertion. It therefore cannot disturb template merges,
        printing layout, row heights, or the blue input-cell formatting.
        """

        first = LAYOUT.first_program_row - 1
        last = LAYOUT.last_program_row - 1
        if direction not in (-1, 1) or not first <= start_table_row <= last:
            return
        if direction == 1:
            source_rows = range(last - 1, start_table_row - 1, -1)
            clear_row = start_table_row
        else:
            source_rows = range(start_table_row + 1, last + 1)
            clear_row = last
        self._loading = True
        try:
            # Column A contains the fixed serial numbers, so only move the
            # user/program data area B onward.
            for column in range(1, self.columnCount()):
                for source_row in source_rows:
                    source = self.item(source_row, column)
                    destination = self.item(source_row + direction, column)
                    if source is not None and destination is not None:
                        destination.setText(source.text())
                item = self.item(clear_row, column)
                if item is not None:
                    item.setText("")
        finally:
            self._loading = False
        self._sync_known_values()

    @staticmethod
    def _theme_palette(theme: bytes | None) -> list[str]:
        if not theme:
            return []
        try:
            root = ElementTree.fromstring(theme)
            namespace = "{http://schemas.openxmlformats.org/drawingml/2006/main}"
            scheme = root.find(f".//{namespace}clrScheme")
            if scheme is None:
                return []
            values: list[str] = []
            for child in list(scheme):
                color = child.find(f"{namespace}srgbClr")
                if color is None:
                    color = child.find(f"{namespace}sysClr")
                if color is not None:
                    values.append(color.attrib.get("val") or color.attrib.get("lastClr") or "000000")
            return values
        except ElementTree.ParseError:
            return []

    def _rgb(self, openpyxl_color: Any) -> QColor | None:
        if openpyxl_color is not None and getattr(openpyxl_color, "type", "") == "rgb":
            value = getattr(openpyxl_color, "rgb", None)
            if value and len(value) >= 6:
                return QColor("#" + value[-6:])
        if openpyxl_color is not None and getattr(openpyxl_color, "type", "") == "theme":
            index = int(getattr(openpyxl_color, "theme", -1))
            if 0 <= index < len(self._theme_colors):
                color = QColor("#" + self._theme_colors[index][-6:])
                tint = float(getattr(openpyxl_color, "tint", 0.0) or 0.0)
                if tint:
                    channels = [color.red(), color.green(), color.blue()]
                    if tint > 0:
                        channels = [round(channel + (255 - channel) * tint) for channel in channels]
                    else:
                        channels = [round(channel * (1 + tint)) for channel in channels]
                    return QColor(*channels)
                return color
        return None

    @staticmethod
    def _editable(cell: Any) -> bool:
        if cell.data_type == "f":
            return False
        # Static template captions are protected in the preview. Blue/form fields,
        # all program rows, and blank header information cells remain editable.
        if cell.row <= 6 and cell.value not in (None, ""):
            return False
        if cell.column == 1 and cell.row >= 5:
            return False
        return True

    def load_template(self, template_path: Path) -> None:
        workbook = load_workbook(template_path, read_only=False, data_only=False, keep_links=True)
        try:
            sheet = workbook[LAYOUT.sheet_name]
            self._loading = True
            self._theme_colors = self._theme_palette(workbook.loaded_theme)
            self._clear_image_overlays()
            self._merged_cell_roots.clear()
            self.clear()
            self.setRowCount(sheet.max_row)
            self.setColumnCount(sheet.max_column)
            self.setHorizontalHeaderLabels([self._column_name(index) for index in range(1, sheet.max_column + 1)])
            self.setVerticalHeaderLabels([str(index) for index in range(1, sheet.max_row + 1)])
            for column in range(1, sheet.max_column + 1):
                width = sheet.column_dimensions[self._column_name(column)].width
                self.setColumnWidth(column - 1, max(44, int((width or 10) * 7 + 10)))
            for row in range(1, sheet.max_row + 1):
                height = sheet.row_dimensions[row].height
                self.setRowHeight(row - 1, max(23, int((height or 15) * 1.35)))
            for row in sheet.iter_rows():
                for cell in row:
                    if isinstance(cell, MergedCell):
                        continue
                    item = QTableWidgetItem("" if cell.value is None else str(cell.value))
                    item.setToolTip(cell.coordinate)
                    # A no-fill cell carries ``00000000`` in openpyxl. That
                    # is transparent/default, not a black Excel fill.
                    fill = self._rgb(cell.fill.fgColor) if cell.fill.fill_type else None
                    item.setBackground(fill or QColor("#FFFFFF"))
                    font_color = self._rgb(cell.font.color)
                    font = QFont(cell.font.name or "Arial", int(cell.font.sz or 10))
                    font.setBold(bool(cell.font.bold))
                    font.setItalic(bool(cell.font.italic))
                    item.setFont(font)
                    item.setForeground(font_color or QColor("#000000"))
                    alignment = Qt.AlignmentFlag.AlignVCenter
                    if cell.alignment.horizontal == "center":
                        alignment |= Qt.AlignmentFlag.AlignHCenter
                    elif cell.alignment.horizontal == "right":
                        alignment |= Qt.AlignmentFlag.AlignRight
                    else:
                        alignment |= Qt.AlignmentFlag.AlignLeft
                    item.setTextAlignment(int(alignment))
                    if not self._editable(cell):
                        item.setFlags(item.flags() & ~Qt.ItemFlag.ItemIsEditable)
                    item.setData(Qt.ItemDataRole.UserRole, cell.coordinate)
                    self.setItem(cell.row - 1, cell.column - 1, item)
            for merged in sheet.merged_cells.ranges:
                self.setSpan(merged.min_row - 1, merged.min_col - 1, merged.max_row - merged.min_row + 1, merged.max_col - merged.min_col + 1)
                root_coordinate = sheet.cell(merged.min_row, merged.min_col).coordinate
                for row in range(merged.min_row - 1, merged.max_row):
                    for column in range(merged.min_col - 1, merged.max_col):
                        self._merged_cell_roots[(row, column)] = root_coordinate
        finally:
            self._loading = False
            workbook.close()
        self._sync_known_values()
        self.clear_edit_history()
        QTimer.singleShot(0, self._reposition_image_overlays)

    @staticmethod
    def _column_name(number: int) -> str:
        name = ""
        while number:
            number, remainder = divmod(number - 1, 26)
            name = chr(65 + remainder) + name
        return name

    def set_value(self, cell: str, value: str) -> None:
        column_letters = "".join(character for character in cell if character.isalpha())
        row_number = int("".join(character for character in cell if character.isdigit()))
        column = 0
        for character in column_letters:
            column = column * 26 + ord(character) - 64
        item = self.item(row_number - 1, column - 1)
        if item is None:
            item = QTableWidgetItem()
            self.setItem(row_number - 1, column - 1, item)
            item.setData(Qt.ItemDataRole.UserRole, cell)
        self._loading = True
        item.setText(value)
        self._loading = False
        self._known_values[cell] = value

    def set_text_color(self, cell: str, color: QColor) -> None:
        """Change a preview cell's text color without replacing its style."""

        column_letters = "".join(character for character in cell if character.isalpha())
        row_number = int("".join(character for character in cell if character.isdigit()))
        column = 0
        for character in column_letters:
            column = column * 26 + ord(character) - 64
        item = self.item(row_number - 1, column - 1)
        if item is not None:
            item.setForeground(color)

    def current_coordinate(self) -> str | None:
        """Return the selected Excel coordinate, including merged-cell roots."""

        row, column = self.currentRow(), self.currentColumn()
        if row < 0 or column < 0:
            return None
        item = self.item(row, column)
        coordinate = item.data(Qt.ItemDataRole.UserRole) if item is not None else None
        if coordinate:
            return str(coordinate)
        return self._merged_cell_roots.get((row, column))

    @property
    def selected_image_id(self) -> str | None:
        return self._selected_image_id

    @staticmethod
    def _column_number(letters: str) -> int:
        number = 0
        for character in letters:
            number = number * 26 + ord(character) - 64
        return number

    def _sheet_size(self) -> tuple[int, int]:
        return (
            sum(self.columnWidth(column) for column in range(self.columnCount())),
            sum(self.rowHeight(row) for row in range(self.rowCount())),
        )

    def _placement_sheet_position(self, placement: ImagePlacement) -> tuple[int, int]:
        letters = "".join(character for character in placement.anchor if character.isalpha())
        digits = "".join(character for character in placement.anchor if character.isdigit())
        if not letters or not digits:
            return placement.offset_x, placement.offset_y
        column = self._column_number(letters) - 1
        row = int(digits) - 1
        x = sum(self.columnWidth(index) for index in range(max(0, column))) + placement.offset_x
        y = sum(self.rowHeight(index) for index in range(max(0, row))) + placement.offset_y
        return x, y

    def _anchor_for_sheet_position(self, x: int, y: int) -> tuple[str, int, int]:
        sheet_width, sheet_height = self._sheet_size()
        x = min(max(0, round(x)), max(0, sheet_width - 1))
        y = min(max(0, round(y)), max(0, sheet_height - 1))
        column = 0
        column_start = 0
        for index in range(self.columnCount()):
            width = self.columnWidth(index)
            if x < column_start + width or index == self.columnCount() - 1:
                column = index
                break
            column_start += width
        row = 0
        row_start = 0
        for index in range(self.rowCount()):
            height = self.rowHeight(index)
            if y < row_start + height or index == self.rowCount() - 1:
                row = index
                break
            row_start += height
        return f"{self._column_name(column + 1)}{row + 1}", x - column_start, y - row_start

    def sheet_to_viewport(self, x: int, y: int) -> tuple[int, int]:
        return x - self.horizontalScrollBar().value(), y - self.verticalScrollBar().value()

    def viewport_to_sheet(self, x: int, y: int) -> tuple[int, int]:
        return x + self.horizontalScrollBar().value(), y + self.verticalScrollBar().value()

    def constrain_image_geometry(
        self, x: int, y: int, width: int, height: int, minimum_size: int = 32
    ) -> tuple[int, int, int, int]:
        sheet_width, sheet_height = self._sheet_size()
        width = min(max(minimum_size, round(width)), max(minimum_size, sheet_width))
        height = min(max(minimum_size, round(height)), max(minimum_size, sheet_height))
        x = min(max(0, round(x)), max(0, sheet_width - width))
        y = min(max(0, round(y)), max(0, sheet_height - height))
        return x, y, width, height

    def _clear_image_overlays(self) -> None:
        for overlay in self._image_overlays.values():
            overlay.hide()
            overlay.deleteLater()
        self._image_overlays.clear()
        self._selected_image_id = None

    def refresh_floating_images(self, placements: list[ImagePlacement]) -> None:
        """Render operator pictures as movable objects above the spreadsheet."""

        wanted_ids = {placement.placement_id for placement in placements}
        for placement_id, overlay in tuple(self._image_overlays.items()):
            if placement_id not in wanted_ids:
                overlay.hide()
                overlay.deleteLater()
                del self._image_overlays[placement_id]
        if self._selected_image_id not in wanted_ids:
            self._selected_image_id = None
        for placement in placements:
            overlay = self._image_overlays.get(placement.placement_id)
            if overlay is None:
                overlay = FloatingImageOverlay(self, placement)
                overlay.selected.connect(self._select_image)
                overlay.geometry_committed.connect(self._commit_overlay_geometry)
                overlay.delete_requested.connect(self.image_delete_requested)
                self._image_overlays[placement.placement_id] = overlay
            else:
                overlay.set_placement(placement)
            overlay.set_selected(placement.placement_id == self._selected_image_id)
        self._reposition_image_overlays()

    def _reposition_image_overlays(self) -> None:
        for overlay in self._image_overlays.values():
            x, y = self._placement_sheet_position(overlay.placement)
            overlay.set_sheet_geometry(x, y, overlay.placement.width, overlay.placement.height)
            overlay.show()
            overlay.raise_()
        if self._selected_image_id and self._selected_image_id in self._image_overlays:
            self._image_overlays[self._selected_image_id].raise_()

    def _select_image(self, placement_id: str) -> None:
        self._selected_image_id = placement_id
        for candidate_id, overlay in self._image_overlays.items():
            overlay.set_selected(candidate_id == placement_id)
        selected = self._image_overlays.get(placement_id)
        if selected is not None:
            selected.raise_()

    def select_image(self, placement_id: str) -> None:
        """Select one floating picture after it has been added programmatically."""

        if placement_id in self._image_overlays:
            self._select_image(placement_id)

    def clear_image_selection(self) -> None:
        """Remove editing handles temporarily, for example while printing."""

        self._selected_image_id = None
        for overlay in self._image_overlays.values():
            overlay.set_selected(False)

    def _commit_overlay_geometry(self, placement_id: str, x: int, y: int, width: int, height: int) -> None:
        anchor, offset_x, offset_y = self._anchor_for_sheet_position(x, y)
        self.image_geometry_changed.emit(placement_id, anchor, offset_x, offset_y, width, height)

    def _on_item_changed(self, item: QTableWidgetItem) -> None:
        coordinate = item.data(Qt.ItemDataRole.UserRole)
        if not self._loading and coordinate:
            coordinate = str(coordinate)
            previous = self._known_values.get(coordinate, "")
            current = item.text()
            self._known_values[coordinate] = current
            self._remember_edit([(coordinate, previous, current)])
            self.cell_edited.emit(coordinate, current)


class SettingsDialog(QDialog):
    """The single, familiar settings entry point for normal users."""

    def __init__(
        self,
        current_repository: str,
        show_save_reminder: bool,
        automatic_update_check: bool,
        confirm_before_clear: bool,
        confirm_before_save: bool,
        open_folder_after_save: bool,
        toolbar_visibility: dict[str, bool],
        print_orientation: str,
        print_margin_mm: int,
        selection_color: str,
        parent: QWidget | None = None,
    ) -> None:
        super().__init__(parent)
        self.setWindowTitle("设置")
        self.resize(540, 375)
        self.repository = QLineEdit(current_repository)
        self.repository.setPlaceholderText("owner/repository")
        self.save_reminder = QCheckBox("保存完成后显示提醒")
        self.save_reminder.setChecked(show_save_reminder)
        self.automatic_update_check = QCheckBox("启动后及后台自动检查更新")
        self.automatic_update_check.setChecked(automatic_update_check)
        self.confirm_before_clear = QCheckBox("清空已导入内容前要求确认")
        self.confirm_before_clear.setChecked(confirm_before_clear)
        self.confirm_before_save = QCheckBox("保存程序单前显示确认窗口")
        self.confirm_before_save.setChecked(confirm_before_save)
        self.open_folder_after_save = QCheckBox("保存完成后打开所在文件夹")
        self.open_folder_after_save.setChecked(open_folder_after_save)

        root = QVBoxLayout(self)
        tabs = QTabWidget()
        root.addWidget(tabs, 1)

        general = QWidget()
        general_layout = QVBoxLayout(general)
        general_layout.addWidget(QLabel(f"{APP_NAME}  v{__version__}"))
        general_note = QLabel("程序模板、保存位置、图片与个人设置均保存在本机 AppData；升级软件不会删除这些内容。")
        general_note.setWordWrap(True)
        general_layout.addWidget(general_note)
        general_layout.addWidget(self.open_folder_after_save)
        general_layout.addStretch(1)
        tabs.addTab(general, "常规")

        editing = QWidget()
        editing_layout = QVBoxLayout(editing)
        editing_layout.addWidget(self.confirm_before_clear)
        editing_layout.addWidget(self.confirm_before_save)
        self.selection_color = QLineEdit(self._normalized_selection_color(selection_color))
        self.selection_color.setMaxLength(7)
        self.selection_color.setPlaceholderText("#B7D7FF")
        self.selection_color.setToolTip("输入 #RRGGBB，或点击“选择颜色”使用调色板")
        self.selection_color.textChanged.connect(self._refresh_selection_color_sample)
        self.selection_color_sample = QLabel("框选示例")
        self.selection_color_sample.setAlignment(Qt.AlignmentFlag.AlignCenter)
        self.selection_color_sample.setMinimumWidth(75)
        choose_color = QPushButton("选择颜色…")
        choose_color.clicked.connect(self._choose_selection_color)
        restore_color = QPushButton("默认")
        restore_color.clicked.connect(lambda: self.selection_color.setText(DEFAULT_SELECTION_COLOR))
        color_row = QHBoxLayout()
        color_row.addWidget(self.selection_color, 1)
        color_row.addWidget(self.selection_color_sample)
        color_row.addWidget(choose_color)
        color_row.addWidget(restore_color)
        editing_layout.addWidget(QLabel("框选颜色："))
        editing_layout.addLayout(color_row)
        self._refresh_selection_color_sample()
        editing_note = QLabel(
            "表格支持 Ctrl+Z 撤销、Ctrl+Y 重做、Ctrl+F 查找、Ctrl+H 替换；"
            "框选后单击空白处或按 Esc 可完全取消框选。"
        )
        editing_note.setWordWrap(True)
        editing_layout.addWidget(editing_note)
        editing_layout.addStretch(1)
        tabs.addTab(editing, "编辑")

        toolbar_page = QWidget()
        toolbar_layout = QVBoxLayout(toolbar_page)
        toolbar_note = QLabel("取消勾选后，按钮会立刻从顶部工具栏隐藏；需要时可随时回到“设置”重新勾选。完成并保存、选择 NC 文件、设置会始终保留。")
        toolbar_note.setWordWrap(True)
        toolbar_layout.addWidget(toolbar_note)
        self.toolbar_buttons: dict[str, QCheckBox] = {}
        check_rows = QHBoxLayout()
        left_checks = QVBoxLayout()
        right_checks = QVBoxLayout()
        for index, (key, label) in enumerate(OPTIONAL_TOOLBAR_BUTTONS):
            checkbox = QCheckBox(label)
            checkbox.setChecked(toolbar_visibility.get(key, True))
            self.toolbar_buttons[key] = checkbox
            (left_checks if index < (len(OPTIONAL_TOOLBAR_BUTTONS) + 1) // 2 else right_checks).addWidget(checkbox)
        check_rows.addLayout(left_checks)
        check_rows.addLayout(right_checks)
        toolbar_layout.addLayout(check_rows)
        toolbar_layout.addStretch(1)
        tabs.addTab(toolbar_page, "工具栏")

        printing = QWidget()
        print_form = QFormLayout(printing)
        self.print_orientation = QComboBox()
        self.print_orientation.addItem("▯ 竖向（默认）", "portrait")
        self.print_orientation.addItem("▭ 横向", "landscape")
        self.print_orientation.setCurrentIndex(0 if print_orientation != "landscape" else 1)
        self.print_margin_mm = QSpinBox()
        self.print_margin_mm.setRange(0, 30)
        self.print_margin_mm.setSuffix(" mm")
        self.print_margin_mm.setValue(max(0, min(30, print_margin_mm)))
        print_form.addRow("默认方向：", self.print_orientation)
        print_form.addRow("页面边距：", self.print_margin_mm)
        print_note = QLabel("打印区域、方向与边距会自动记住。可在打印预览底部点击“打印区域”回到主表格框选，也可右键设为打印区域。")
        print_note.setWordWrap(True)
        print_form.addRow(print_note)
        tabs.addTab(printing, "打印")

        reminders = QWidget()
        reminder_layout = QVBoxLayout(reminders)
        reminder_layout.addWidget(self.save_reminder)
        reminder_note = QLabel("关闭后，保存成功只在窗口底部显示状态，不再弹出提示。随时可在这里重新打开。")
        reminder_note.setWordWrap(True)
        reminder_layout.addWidget(reminder_note)
        reminder_layout.addStretch(1)
        tabs.addTab(reminders, "提醒")

        updates = QWidget()
        update_form = QFormLayout(updates)
        update_form.addRow(self.automatic_update_check)
        update_form.addRow("GitHub 更新仓库：", self.repository)
        update_note = QLabel("正常用户无需改动仓库。关闭自动检查后，仍可随时点击主界面的“检查更新”。")
        update_note.setWordWrap(True)
        update_form.addRow(update_note)
        tabs.addTab(updates, "更新")

        buttons = QDialogButtonBox(QDialogButtonBox.StandardButton.Save | QDialogButtonBox.StandardButton.Cancel)
        buttons.accepted.connect(self.accept)
        buttons.rejected.connect(self.reject)
        root.addWidget(buttons)

    @staticmethod
    def _normalized_selection_color(color_name: str) -> str:
        color = QColor(color_name)
        return color.name() if color.isValid() else DEFAULT_SELECTION_COLOR

    def selected_selection_color(self) -> str:
        """Return a safe persisted value even if the typed text is invalid."""

        return self._normalized_selection_color(self.selection_color.text().strip())

    def _refresh_selection_color_sample(self) -> None:
        color = self.selected_selection_color()
        self.selection_color_sample.setStyleSheet(
            f"background-color: {color}; color: #000000; border: 1px solid {QColor(color).darker(175).name()};"
        )

    def _choose_selection_color(self) -> None:
        color = QColorDialog.getColor(QColor(self.selected_selection_color()), self, "选择框选颜色")
        if color.isValid():
            self.selection_color.setText(color.name())


class FindReplaceDialog(QDialog):
    """Small Excel-like find/replace panel for the editable preview sheet."""

    def __init__(self, preview: ExcelPreview, replace_mode: bool = True, parent: QWidget | None = None) -> None:
        super().__init__(parent)
        self.preview = preview
        self.setWindowTitle("查找和替换")
        self.setMinimumWidth(410)
        self.find_text = QLineEdit()
        self.find_text.setPlaceholderText("输入要查找的内容")
        self.replace_text = QLineEdit()
        self.replace_text.setPlaceholderText("输入替换后的内容")
        self.match_case = QCheckBox("区分大小写")
        self.status = QLabel("在右侧表格的可编辑单元格中查找。")
        self.status.setWordWrap(True)

        form = QFormLayout()
        form.addRow("查找内容：", self.find_text)
        form.addRow("替换为：", self.replace_text)
        form.addRow("", self.match_case)
        root = QVBoxLayout(self)
        root.addLayout(form)
        root.addWidget(self.status)
        actions = QHBoxLayout()
        find_next = QPushButton("查找下一个")
        find_next.clicked.connect(self.find_next)
        actions.addWidget(find_next)
        replace = QPushButton("替换")
        replace.clicked.connect(self.replace_current)
        actions.addWidget(replace)
        replace_all = QPushButton("全部替换")
        replace_all.clicked.connect(self.replace_all)
        actions.addWidget(replace_all)
        actions.addStretch(1)
        close = QPushButton("关闭")
        close.clicked.connect(self.close)
        actions.addWidget(close)
        root.addLayout(actions)
        if not replace_mode:
            self.replace_text.setFocus()
            self.find_text.setFocus()
        else:
            self.find_text.setFocus()

    def _find_value(self) -> str:
        return self.find_text.text()

    def find_next(self) -> None:
        text = self._find_value()
        if not text:
            self.status.setText("请先输入要查找的内容。")
            self.find_text.setFocus()
            return
        coordinate = self.preview.find_next(text, self.match_case.isChecked())
        self.status.setText(f"已定位到 {coordinate}。" if coordinate else f"没有找到“{text}”。")

    def replace_current(self) -> None:
        text = self._find_value()
        if not text:
            self.find_next()
            return
        if self.preview.replace_current(text, self.replace_text.text(), self.match_case.isChecked()):
            self.status.setText("已替换当前单元格，可按 Ctrl+Z 撤销。")
        else:
            self.status.setText("当前单元格没有可替换的内容，已继续查找。")
            self.find_next()

    def replace_all(self) -> None:
        text = self._find_value()
        if not text:
            self.find_next()
            return
        count = self.preview.replace_all(text, self.replace_text.text(), self.match_case.isChecked())
        self.status.setText(f"已在 {count} 个单元格中完成替换，可按 Ctrl+Z 撤销。" if count else f"没有找到“{text}”。")


class SaveDirectoryDialog(QDialog):
    """Choose the persistent folder used by the save button."""

    def __init__(self, current_directory: Path, parent: QWidget | None = None) -> None:
        super().__init__(parent)
        self.setWindowTitle("设置程序单保存位置")
        self.directory = QLineEdit(str(current_directory))
        browse = QPushButton("选择文件夹")
        browse.clicked.connect(self._browse)
        row = QHBoxLayout()
        row.addWidget(self.directory, 1)
        row.addWidget(browse)
        form = QFormLayout(self)
        form.addRow("固定保存文件夹：", row)
        note = QLabel(
            "首次保存后会记住此位置。以后每次点击“完成并保存”都会保存到这里；"
            "在模板顶部右侧空白名称框输入文字后，导出文件会使用“名称_日期.xlsx”。"
        )
        note.setWordWrap(True)
        form.addRow(note)
        buttons = QDialogButtonBox(QDialogButtonBox.StandardButton.Save | QDialogButtonBox.StandardButton.Cancel)
        buttons.accepted.connect(self.accept)
        buttons.rejected.connect(self.reject)
        form.addRow(buttons)

    def _browse(self) -> None:
        selected = QFileDialog.getExistingDirectory(self, "选择程序单保存文件夹", self.directory.text())
        if selected:
            self.directory.setText(selected)


class WheelZoomPrintPreview(QPrintPreviewWidget):
    """Preview widget whose mouse wheel always changes the page scale."""

    def wheelEvent(self, event: Any) -> None:
        delta = event.angleDelta().y()
        if not delta:
            super().wheelEvent(event)
            return
        factor = self.zoomFactor() or 1.0
        multiplier = 1.15 if delta > 0 else 1 / 1.15
        self.setZoomFactor(max(0.15, min(8.0, factor * multiplier)))
        event.accept()


class PrinterSelectionDialog(QDialog):
    """Choose a Windows printer without opening the unsupported OS preview."""

    def __init__(
        self,
        orientation: QPageLayout.Orientation = QPageLayout.Orientation.Landscape,
        parent: QWidget | None = None,
    ) -> None:
        super().__init__(parent)
        self.setWindowTitle("选择打印机")
        self.setMinimumWidth(380)
        self.printer_names = QComboBox()
        available = QPrinterInfo.availablePrinterNames()
        default_name = QPrinterInfo.defaultPrinter().printerName()
        self.printer_names.addItems(available)
        if default_name:
            index = self.printer_names.findText(default_name)
            if index >= 0:
                self.printer_names.setCurrentIndex(index)
        self.copies = QSpinBox()
        self.copies.setRange(1, 99)
        self.copies.setValue(1)
        self.orientation = QComboBox()
        # The page-shaped symbols make the choice recognisable at a glance,
        # while retaining the Chinese label for unambiguous operation.
        self.orientation.addItem("▯ 竖向（Portrait）", QPageLayout.Orientation.Portrait)
        self.orientation.addItem("▭ 横向（Landscape）", QPageLayout.Orientation.Landscape)
        orientation_index = self.orientation.findData(orientation)
        self.orientation.setCurrentIndex(orientation_index if orientation_index >= 0 else 1)

        form = QFormLayout(self)
        form.addRow("打印机：", self.printer_names)
        form.addRow("方向：", self.orientation)
        form.addRow("份数：", self.copies)
        note = QLabel("打印前可在软件预览中确认方向。此窗口直接发送 Windows 打印任务，不使用 Windows 不支持预览的系统打印窗口。")
        note.setWordWrap(True)
        form.addRow(note)
        buttons = QDialogButtonBox(QDialogButtonBox.StandardButton.Ok | QDialogButtonBox.StandardButton.Cancel)
        buttons.button(QDialogButtonBox.StandardButton.Ok).setText("打印")
        buttons.accepted.connect(self.accept)
        buttons.rejected.connect(self.reject)
        form.addRow(buttons)

    def build_printer(self) -> QPrinter | None:
        name = self.printer_names.currentText().strip()
        if not name:
            return None
        printer = QPrinter(QPrinter.PrinterMode.HighResolution)
        printer.setPrinterName(name)
        printer.setPageSize(QPageSize(QPageSize.PageSizeId.A4))
        printer.setPageOrientation(self.orientation.currentData())
        printer.setCopyCount(self.copies.value())
        return printer


class ProgramSheetPrintPreviewDialog(QDialog):
    """Printable sheet preview with explicit controls and wheel zoom."""

    select_print_area_requested = Signal()

    def __init__(
        self,
        printer: QPrinter,
        render_page: Callable[[QPrinter], None],
        print_area_label: str,
        margin_mm: int,
        preferences_changed: Callable[[QPageLayout.Orientation, int], None] | None = None,
        parent: QWidget | None = None,
    ) -> None:
        super().__init__(parent)
        self._printer = printer
        self._render_page = render_page
        self._preferences_changed = preferences_changed
        self.setWindowTitle("打印预览 - CNC 程序单")
        self.resize(1020, 760)
        root = QVBoxLayout(self)
        note = QLabel(f"纸张：A4　·　打印区域：{print_area_label}　·　鼠标滚轮可放大/缩小")
        note.setStyleSheet("color: #335b7d;")
        root.addWidget(note)
        self.preview = WheelZoomPrintPreview(printer, self)
        self.preview.setZoomMode(QPrintPreviewWidget.ZoomMode.FitInView)
        self.preview.paintRequested.connect(self._render_page)
        self.preview.previewChanged.connect(self._update_zoom_label)
        root.addWidget(self.preview, 1)
        controls = QHBoxLayout()
        zoom_out = QPushButton("－")
        zoom_out.setToolTip("缩小")
        zoom_out.clicked.connect(lambda: self._change_zoom(1 / 1.15))
        controls.addWidget(zoom_out)
        zoom_in = QPushButton("＋")
        zoom_in.setToolTip("放大")
        zoom_in.clicked.connect(lambda: self._change_zoom(1.15))
        controls.addWidget(zoom_in)
        self.zoom_label = QLabel()
        controls.addWidget(self.zoom_label)
        fit_page = QPushButton("适应页面")
        fit_page.clicked.connect(lambda: self.preview.setZoomMode(QPrintPreviewWidget.ZoomMode.FitInView))
        controls.addWidget(fit_page)
        fit_width = QPushButton("适应宽度")
        fit_width.clicked.connect(lambda: self.preview.setZoomMode(QPrintPreviewWidget.ZoomMode.FitToWidth))
        controls.addWidget(fit_width)
        self.orientation = QComboBox()
        self.orientation.addItem("▯ 竖向", QPageLayout.Orientation.Portrait)
        self.orientation.addItem("▭ 横向", QPageLayout.Orientation.Landscape)
        current_orientation = self._printer.pageLayout().orientation()
        self.orientation.setCurrentIndex(max(0, self.orientation.findData(current_orientation)))
        self.orientation.setToolTip("选择纸张方向，预览会立即更新")
        self.orientation.currentIndexChanged.connect(self._change_orientation)
        controls.addWidget(self.orientation)
        controls.addWidget(QLabel("边距："))
        self.margin = QSpinBox()
        self.margin.setRange(0, 30)
        self.margin.setSuffix(" mm")
        self.margin.setValue(margin_mm)
        self.margin.setToolTip("调整打印区域距离纸张边缘的距离")
        self.margin.valueChanged.connect(self._change_margin)
        controls.addWidget(self.margin)
        print_area = QPushButton("打印区域")
        print_area.setToolTip("返回主表格框选区域；松开鼠标后会自动记住")
        print_area.clicked.connect(self._select_print_area)
        controls.addWidget(print_area)
        controls.addStretch(1)
        print_button = QPushButton("打印")
        print_button.setObjectName("primaryButton")
        print_button.clicked.connect(self._print)
        controls.addWidget(print_button)
        close_button = QPushButton("关闭")
        close_button.clicked.connect(self.reject)
        controls.addWidget(close_button)
        root.addLayout(controls)
        self._update_zoom_label()

    def _change_zoom(self, multiplier: float) -> None:
        factor = self.preview.zoomFactor() or 1.0
        self.preview.setZoomFactor(max(0.15, min(8.0, factor * multiplier)))
        self._update_zoom_label()

    def _update_zoom_label(self) -> None:
        self.zoom_label.setText(f"缩放：{round((self.preview.zoomFactor() or 1.0) * 100)}%")

    def _change_orientation(self) -> None:
        orientation = self.orientation.currentData()
        if orientation is None:
            return
        self._printer.setPageOrientation(orientation)
        self._save_preferences()
        self.preview.updatePreview()

    def _change_margin(self) -> None:
        self._save_preferences()
        self.preview.updatePreview()

    def _save_preferences(self) -> None:
        if self._preferences_changed is not None:
            self._preferences_changed(self._printer.pageLayout().orientation(), self.margin.value())

    def _select_print_area(self) -> None:
        """Close the preview and hand focus back to the sheet for range picking."""

        self.reject()
        self.select_print_area_requested.emit()

    def _print(self) -> None:
        dialog = PrinterSelectionDialog(self._printer.pageLayout().orientation(), self)
        if dialog.exec() != QDialog.DialogCode.Accepted:
            return
        printer = dialog.build_printer()
        if printer is None:
            QMessageBox.warning(self, "未找到打印机", "Windows 当前没有可用打印机，请先安装或连接打印机。")
            return
        self._printer.setPageOrientation(printer.pageLayout().orientation())
        self._save_preferences()
        self._render_page(printer)



class MainWindow(QMainWindow):
    def __init__(self) -> None:
        super().__init__()
        application = QApplication.instance()
        if application is not None:
            # Tests and small support launchers sometimes construct the
            # window directly instead of calling run(); keep their visual
            # baseline identical to the normal desktop entry point.
            apply_application_theme(application)
        self.settings = load_settings()
        # ``None`` is an intentional visual spacer inserted by the operator.
        # Keeping it in the row list means programs below the spacer retain
        # their exact Excel row when the workbook is exported.
        self.records: list[ProgramRecord | None] = []
        # A marker takes a normal program row, shifts following data down and
        # is exported as red text in column B. It stays separate from parsed
        # records so it cannot be mistaken for an NC source during reparse.
        self.row_markers: dict[int, str] = {}
        # UI-only labels showing which legacy shortcut modified a source. They
        # never alter the program name written into the Excel template.
        self.conversion_tags: dict[Path, str] = {}
        # The M06 shortcut intentionally removes the controller's tool-change
        # boundaries. Keep the reviewed pre-conversion text only for this
        # session so a multi-tool program still has one program-sheet row per
        # tool after the exact source-file conversion.
        self.m06_parse_snapshots: dict[Path, str] = {}
        self.manual_cells: dict[str, str] = {}
        self.image_placements: list[ImagePlacement] = []
        self.template_path: Path | None = None
        self._program_area_reset = False
        self._result_loading = False
        self._update_worker: Worker | None = None
        self._print_printer: QPrinter | None = None
        self._awaiting_print_area_selection = False
        # An update installer must be allowed to close the running process
        # without the ordinary user-facing close/minimize question.
        self._exiting_for_update = False
        self.setWindowTitle(f"{APP_NAME}  v{__version__}")
        icon = QIcon(str(resource_path("assets", APP_ICON_FILE)))
        if not icon.isNull():
            self.setWindowIcon(icon)
        self._setup_system_tray(icon)
        self.setMinimumSize(1180, 720)
        self.setAcceptDrops(True)
        self._build_ui()
        self._load_initial_template()
        QTimer.singleShot(1200, self._automatic_update_check)
        self._background_update_timer = QTimer(self)
        self._background_update_timer.setInterval(6 * 60 * 60 * 1000)
        self._background_update_timer.timeout.connect(self._automatic_update_check)
        self._background_update_timer.start()

    def _setup_system_tray(self, icon: QIcon) -> None:
        """Provide the Windows lower-right notification-area application icon."""

        self.tray_icon = QSystemTrayIcon(icon, self)
        self.tray_icon.setToolTip(APP_NAME)
        menu = QMenu(self)
        restore = QAction("显示主窗口", self)
        restore.triggered.connect(self._restore_from_tray)
        quit_action = QAction("退出软件", self)
        quit_action.triggered.connect(self._quit_from_tray)
        menu.addAction(restore)
        menu.addSeparator()
        menu.addAction(quit_action)
        self.tray_icon.setContextMenu(menu)
        self.tray_icon.activated.connect(self._on_tray_activated)
        if QSystemTrayIcon.isSystemTrayAvailable():
            self.tray_icon.show()

    def _hide_to_tray(self) -> None:
        """Hide the main window and leave only its lower-right tray icon."""

        if self.tray_icon.isVisible():
            self.hide()
            self.tray_icon.showMessage(
                APP_NAME,
                "软件已缩小到电脑右下角隐藏图标区。双击图标可恢复窗口。",
                QSystemTrayIcon.MessageIcon.Information,
                2500,
            )
        else:
            # A notification area is normally present on Windows. Keep a
            # usable fallback if a special shell session does not expose one.
            self.showMinimized()

    def _restore_from_tray(self) -> None:
        self.showNormal()
        self.raise_()
        self.activateWindow()

    def _quit_from_tray(self) -> None:
        self._exiting_for_update = True
        self.tray_icon.hide()
        QApplication.instance().quit()

    def _on_tray_activated(self, reason: QSystemTrayIcon.ActivationReason) -> None:
        if reason in {QSystemTrayIcon.ActivationReason.Trigger, QSystemTrayIcon.ActivationReason.DoubleClick}:
            self._restore_from_tray()

    def _build_ui(self) -> None:
        root = QWidget(self)
        layout = QVBoxLayout(root)
        layout.setContentsMargins(12, 12, 12, 12)
        toolbar = QHBoxLayout()
        self._optional_toolbar_buttons: dict[str, QPushButton] = {}
        self.save_button = QPushButton("完成并保存")
        self.save_button.setObjectName("primaryButton")
        self.save_button.clicked.connect(self.save_output)
        toolbar.addWidget(self.save_button)
        save_location = QPushButton("保存位置")
        save_location.clicked.connect(self.configure_save_directory)
        save_location.setToolTip("设置“完成并保存”固定使用的文件夹")
        toolbar.addWidget(save_location)
        reparse = QPushButton("重新解析")
        reparse.clicked.connect(self.reparse)
        toolbar.addWidget(reparse)
        undo = QPushButton("撤销")
        undo.clicked.connect(self.undo_preview_edit)
        undo.setToolTip("撤销表格最后一次编辑（Ctrl+Z）")
        toolbar.addWidget(undo)
        redo = QPushButton("重做")
        redo.clicked.connect(self.redo_preview_edit)
        redo.setToolTip("重做表格最后一次撤销（Ctrl+Y）")
        toolbar.addWidget(redo)
        find_replace = QPushButton("查找替换")
        find_replace.clicked.connect(lambda: self.open_find_replace(True))
        find_replace.setToolTip("在程序单中查找或批量替换内容（Ctrl+F / Ctrl+H）")
        toolbar.addWidget(find_replace)
        clear = QPushButton("清空")
        clear.clicked.connect(self.clear_all)
        toolbar.addWidget(clear)
        template = QPushButton("打开模板")
        template.clicked.connect(self.open_template)
        toolbar.addWidget(template)
        print_sheet = QPushButton("打印")
        print_sheet.clicked.connect(self.preview_program_sheet)
        print_sheet.setToolTip("先预览纸张方向、分页与缩放效果，再选择打印机")
        toolbar.addWidget(print_sheet)
        m06 = QPushButton("M06")
        m06.clicked.connect(lambda: self.open_conversion_dialog(TransformMode.M06))
        m06.setToolTip("按原 M6 快捷方式修改 NC 程序，并自动填写程序单")
        toolbar.addWidget(m06)
        m304 = QPushButton("M304")
        m304.clicked.connect(lambda: self.open_conversion_dialog(TransformMode.M304))
        m304.setToolTip("按原 M304 快捷方式修改 NC 程序，并自动填写程序单")
        toolbar.addWidget(m304)
        copper = QPushButton("铜工")
        copper.clicked.connect(lambda: self.open_conversion_dialog(TransformMode.COPPER))
        copper.setToolTip("按原铜工快捷方式修改 NC 程序，并自动填写程序单")
        toolbar.addWidget(copper)
        add_files = QPushButton("选择 NC 文件")
        add_files.clicked.connect(self.choose_nc_files)
        toolbar.addWidget(add_files)
        toolbar.addStretch(1)
        updates = QPushButton("检查更新")
        updates.clicked.connect(lambda: self.check_updates(silent=False))
        toolbar.addWidget(updates)
        settings_button = QPushButton("设置")
        settings_button.clicked.connect(self.edit_settings)
        settings_button.setToolTip("常规、提醒与更新设置")
        toolbar.addWidget(settings_button)
        self._optional_toolbar_buttons = {
            "save_location": save_location,
            "reparse": reparse,
            "undo": undo,
            "redo": redo,
            "find_replace": find_replace,
            "clear": clear,
            "template": template,
            "print": print_sheet,
            "m06": m06,
            "m304": m304,
            "copper": copper,
            "updates": updates,
        }
        self._apply_toolbar_button_visibility()
        layout.addLayout(toolbar)

        self.drop_label = NcConversionDropArea(
            is_supported_nc_file,
            object_name="dropArea",
            instruction=(
                "将 NC 程序拖到这里，或复制后在窗口按 Ctrl+V / 右键粘贴导入\n"
                "支持一次导入多个 .NC / .nc 文件及 NC 程序文本，将按程序名数字从小到大排列"
            ),
            minimum_height=95,
        )
        self.drop_label.paths_dropped.connect(self.add_nc_files)
        self.drop_label.paste_requested.connect(self._paste_nc_from_menu)
        layout.addWidget(self.drop_label)

        splitter = QSplitter(Qt.Orientation.Horizontal)
        left = QWidget()
        left_layout = QVBoxLayout(left)
        left_layout.setContentsMargins(0, 0, 0, 0)
        left_layout.addWidget(QLabel("已导入 NC 文件（双击可查看和手动修改程序）"))
        self.file_list = QListWidget()
        self.file_list.setMinimumWidth(265)
        self.file_list.itemDoubleClicked.connect(self.edit_imported_nc_file)
        self.file_list.setContextMenuPolicy(Qt.ContextMenuPolicy.CustomContextMenu)
        self.file_list.customContextMenuRequested.connect(self._show_nc_import_menu)
        left_layout.addWidget(self.file_list, 2)
        left_layout.addWidget(QLabel("解析结果（可直接修改）"))
        self.result_table = QTableWidget(0, len(RESULT_COLUMNS))
        self.result_table.setHorizontalHeaderLabels([title for title, _ in RESULT_COLUMNS])
        self.result_table.setAlternatingRowColors(True)
        self.result_table.setEditTriggers(QTableWidget.EditTrigger.DoubleClicked | QTableWidget.EditTrigger.EditKeyPressed)
        self.result_table.itemChanged.connect(self._on_result_changed)
        self.result_table.setMinimumWidth(610)
        self.result_table.horizontalHeader().setStretchLastSection(True)
        left_layout.addWidget(self.result_table, 3)

        right = QWidget()
        right_layout = QVBoxLayout(right)
        right_layout.setContentsMargins(0, 0, 0, 0)
        self.template_label = QLabel("Excel 程序单预览")
        right_layout.addWidget(self.template_label)
        self.preview = ExcelPreview()
        self.preview.setObjectName("excelPreview")
        self.preview.set_selection_color(str(self.settings.get("selection_color", DEFAULT_SELECTION_COLOR)))
        self.preview.cell_edited.connect(self._on_preview_edited)
        self.preview.insert_blank_row_requested.connect(self.insert_blank_program_row)
        self.preview.insert_marker_row_requested.connect(self.insert_marker_program_row)
        self.preview.remove_blank_row_requested.connect(self.remove_blank_program_row)
        self.preview.delete_cells_requested.connect(self.delete_preview_cells)
        self.preview.clipboard_image_paste_requested.connect(self.paste_image_from_clipboard)
        self.preview.clipboard_nc_paste_requested.connect(self.paste_nc_from_clipboard)
        self.preview.clipboard_table_paste_requested.connect(self.paste_table_from_clipboard)
        self.preview.image_geometry_changed.connect(self._update_image_geometry)
        self.preview.image_delete_requested.connect(self._remove_image_by_id)
        self.preview.find_replace_requested.connect(self.open_find_replace)
        self.preview.print_area_requested.connect(self.set_print_area)
        self.preview.clear_print_area_requested.connect(self.clear_print_area)
        self.preview.print_area_selection_finished.connect(self._finish_print_area_selection)
        # Every visible work area accepts an NC drop.  Some Qt item views own
        # their viewport and otherwise swallow the drag before it reaches the
        # main window, which made drops appear to do nothing after UI changes.
        self._nc_drop_targets = (
            root,
            splitter,
            left,
            right,
            self.drop_label,
            self.file_list,
            self.file_list.viewport(),
            self.result_table,
            self.result_table.viewport(),
            self.preview,
            self.preview.viewport(),
        )
        # Child widgets normally consume Ctrl+V themselves. Intercept it
        # before an Excel/result cell can interpret copied NC code as plain
        # text, while leaving ordinary text paste untouched for editing.
        for widget in self._nc_drop_targets:
            widget.setAcceptDrops(True)
            widget.installEventFilter(self)
        right_layout.addWidget(self.preview)
        splitter.addWidget(left)
        splitter.addWidget(right)
        splitter.setSizes([560, 820])
        layout.addWidget(splitter, 1)
        self.setCentralWidget(root)
        self.setStatusBar(QStatusBar())

    def _load_initial_template(self) -> None:
        candidate = Path(str(self.settings.get("template_path", "")))
        if not candidate.is_file():
            candidate = default_user_template()
        self.load_template(candidate)

    def load_template(self, path: Path) -> None:
        try:
            verify_template(path)
            self.template_path = path
            self.preview.load_template(path)
            self.image_placements.clear()
            self.preview.refresh_floating_images(self.image_placements)
            self._program_area_reset = False
            # The template's DATE value is a generated field, not a permanent
            # part of the source workbook. Show today's date immediately.
            self.manual_cells.update(automatic_cells_for_rows([]))
            self.preview.set_value(LAYOUT.date_cell, self.manual_cells[LAYOUT.date_cell])
            self.template_label.setText(f"Excel 程序单预览：{path.name}")
            self.settings["template_path"] = str(path)
            save_settings(self.settings)
            self.statusBar().showMessage("模板已加载；原始模板只读使用，不会被覆盖。", 5000)
        except Exception as exc:
            self.show_error("无法打开模板", str(exc))

    def open_template(self) -> None:
        filename, _ = QFileDialog.getOpenFileName(self, "选择 CNC 程序单模板", str(self.template_path or Path.home()), "Excel 文件 (*.xlsx)")
        if filename:
            self.load_template(Path(filename))

    def insert_image(self) -> None:
        """Add a file image as a movable object at the selected Excel cell."""

        anchor = self.preview.current_coordinate()
        if not anchor:
            self.show_error("请选择插入位置", "请先在右侧 Excel 预览中点击一个单元格，再点击“插入图片”。")
            return
        filename, _ = QFileDialog.getOpenFileName(
            self,
            "选择要插入的图片",
            str(Path.home()),
            "图片文件 (*.png *.jpg *.jpeg *.bmp *.gif)",
        )
        if not filename:
            return
        try:
            image_path = Path(filename)
            cached_image = copy_user_image(image_path)
            self._add_image_placement(cached_image, anchor, *self._default_image_size(cached_image))
        except Exception as exc:
            self.show_error("无法插入图片", str(exc))

    @staticmethod
    def _default_image_size(image_path: Path) -> tuple[int, int]:
        """Use a comfortably visible initial size; users can then resize on canvas."""

        image = QImage(str(image_path))
        if image.isNull():
            raise ValueError("无法读取图片文件")
        scale = min(1.0, 520 / image.width(), 360 / image.height())
        # Tiny source files should not turn back into the almost invisible
        # cell icon behaviour this floating editor replaces.  Keep them large
        # enough to grab, without exceeding the initial canvas bounds.
        if max(image.width() * scale, image.height() * scale) < 160:
            scale = min(160 / max(image.width(), image.height()), 520 / image.width(), 360 / image.height())
        return max(1, round(image.width() * scale)), max(1, round(image.height() * scale))

    def _add_image_placement(self, image_path: Path, anchor: str, width: int, height: int) -> None:
        placement = ImagePlacement(source_path=image_path, anchor=anchor, width=width, height=height)
        self.image_placements.append(placement)
        self.preview.refresh_floating_images(self.image_placements)
        self.preview.select_image(placement.placement_id)
        self.statusBar().showMessage("图片已插入：可直接拖动图片移动，拖动蓝色控制点调整大小。", 5000)

    def _update_image_geometry(
        self, placement_id: str, anchor: str, offset_x: int, offset_y: int, width: int, height: int
    ) -> None:
        self.image_placements = [
            replace(
                placement,
                anchor=anchor,
                offset_x=offset_x,
                offset_y=offset_y,
                width=width,
                height=height,
            )
            if placement.placement_id == placement_id
            else placement
            for placement in self.image_placements
        ]
        self.preview.refresh_floating_images(self.image_placements)
        self.preview.select_image(placement_id)
        self.statusBar().showMessage("图片位置和大小已更新；保存后将同步到 Excel。", 3500)

    def paste_image_from_clipboard(self) -> None:
        """Paste a screenshot, copied picture, or copied local image file."""

        anchor = self.preview.current_coordinate()
        if not anchor:
            self.show_error("请选择粘贴位置", "请先在右侧 Excel 预览中点击一个单元格，再按 Ctrl+V。")
            return
        try:
            mime_data = QApplication.clipboard().mimeData()
            image = QApplication.clipboard().image()
            if not image.isNull():
                buffer = QBuffer()
                if not buffer.open(QIODevice.OpenModeFlag.WriteOnly) or not image.save(buffer, "PNG"):
                    raise ValueError("无法读取剪贴板中的图片")
                cached_image = cache_user_image_bytes(bytes(buffer.data()), "clipboard.png")
            else:
                source = next(
                    (
                        Path(url.toLocalFile())
                        for url in mime_data.urls()
                        if url.isLocalFile() and Path(url.toLocalFile()).suffix.lower() in IMAGE_FILE_SUFFIXES
                    ),
                    None,
                )
                if source is None:
                    raise ValueError("剪贴板中没有可粘贴的图片或截图")
                cached_image = copy_user_image(source)
            self._add_image_placement(cached_image, anchor, *self._default_image_size(cached_image))
        except Exception as exc:
            self.show_error("无法粘贴图片", str(exc))

    def paste_table_from_clipboard(self) -> None:
        """Paste the active Win+V item into the selected editable table cell."""

        mime_data = QApplication.clipboard().mimeData()
        image_file_in_clipboard = any(
            url.isLocalFile() and Path(url.toLocalFile()).suffix.lower() in IMAGE_FILE_SUFFIXES
            for url in mime_data.urls()
        )
        if mime_data.hasImage() or image_file_in_clipboard:
            self.paste_image_from_clipboard()
            return
        if mime_data.hasText():
            count = self.preview.paste_text_grid(mime_data.text())
            self.statusBar().showMessage(
                f"已粘贴 {count} 个单元格。" if count else "请选择可编辑单元格后再粘贴。",
                3500,
            )
            return
        self.statusBar().showMessage("当前剪贴板不是可粘贴的文字或图片。请先按 Win+V 选择内容。", 4500)

    def undo_preview_edit(self) -> None:
        if self.preview.undo_edit():
            self.statusBar().showMessage("已撤销表格编辑。", 2500)
        else:
            self.statusBar().showMessage("没有可撤销的表格编辑。", 2500)

    def redo_preview_edit(self) -> None:
        if self.preview.redo_edit():
            self.statusBar().showMessage("已重做表格编辑。", 2500)
        else:
            self.statusBar().showMessage("没有可重做的表格编辑。", 2500)

    def _print_area_key(self) -> str:
        return str(self.template_path.resolve()) if self.template_path is not None else "__default_template__"

    def _full_print_area(self) -> tuple[int, int, int, int]:
        return (0, 0, max(0, self.preview.rowCount() - 1), max(0, self.preview.columnCount() - 1))

    @staticmethod
    def _print_area_label(area: tuple[int, int, int, int]) -> str:
        top, left, bottom, right = area
        start = ExcelPreview._column_name(left + 1) + str(top + 1)
        end = ExcelPreview._column_name(right + 1) + str(bottom + 1)
        return f"{start}:{end}"

    def _saved_print_area(self) -> tuple[int, int, int, int]:
        full = self._full_print_area()
        saved = self.settings.get("print_areas", {})
        raw = saved.get(self._print_area_key()) if isinstance(saved, dict) else None
        if not isinstance(raw, list) or len(raw) != 4 or not all(isinstance(value, int) for value in raw):
            return full
        top, left, bottom, right = raw
        if top < 0 or left < 0 or bottom < top or right < left or bottom >= self.preview.rowCount() or right >= self.preview.columnCount():
            return full
        return top, left, bottom, right

    def set_print_area_from_selection(self) -> None:
        area = self.preview.selected_cell_range()
        if area is None:
            self.show_error("请选择打印区域", "请先在右侧表格中拖动框选要打印的单元格，再点击“打印区域”。")
            return
        self.set_print_area(*area)

    def begin_print_area_selection(self) -> None:
        """Return from print preview and let the user drag a remembered range."""

        self._awaiting_print_area_selection = True
        self.showNormal()
        self.raise_()
        self.activateWindow()
        self.preview.clearSelection()
        self.preview.clear_image_selection()
        self.preview.setFocus()
        self.statusBar().showMessage(
            "请在右侧表格拖动框选打印区域；松开鼠标后会自动记住，再点击“打印”预览。",
            9000,
        )

    def _finish_print_area_selection(self, top: int, left: int, bottom: int, right: int) -> None:
        if not self._awaiting_print_area_selection:
            return
        self._awaiting_print_area_selection = False
        self.set_print_area(top, left, bottom, right)

    def set_print_area(self, top: int, left: int, bottom: int, right: int) -> None:
        """Persist a rectangular Excel range for this template's future prints."""

        full = self._full_print_area()
        top, left = max(0, top), max(0, left)
        bottom, right = min(full[2], bottom), min(full[3], right)
        if bottom < top or right < left:
            self.show_error("无法设置打印区域", "当前框选区域无效。")
            return
        areas = self.settings.get("print_areas", {})
        if not isinstance(areas, dict):
            areas = {}
        area = (top, left, bottom, right)
        areas[self._print_area_key()] = list(area)
        self.settings["print_areas"] = areas
        save_settings(self.settings)
        self.statusBar().showMessage(f"已记住打印区域：{self._print_area_label(area)}", 4000)

    def clear_print_area(self) -> None:
        areas = self.settings.get("print_areas", {})
        if isinstance(areas, dict):
            areas.pop(self._print_area_key(), None)
            self.settings["print_areas"] = areas
            save_settings(self.settings)
        self.statusBar().showMessage("已清除打印区域，将打印整张表。", 3500)

    def _configured_print_orientation(self) -> QPageLayout.Orientation:
        return (
            QPageLayout.Orientation.Landscape
            if self.settings.get("print_orientation") == "landscape"
            else QPageLayout.Orientation.Portrait
        )

    def _configured_print_margin_mm(self) -> int:
        try:
            return max(0, min(30, int(self.settings.get("print_margin_mm", 8))))
        except (TypeError, ValueError):
            return 8

    def _save_print_preferences(self, orientation: QPageLayout.Orientation, margin_mm: int) -> None:
        self.settings["print_orientation"] = "landscape" if orientation == QPageLayout.Orientation.Landscape else "portrait"
        self.settings["print_margin_mm"] = max(0, min(30, int(margin_mm)))
        save_settings(self.settings)

    def _capture_sheet_for_printing(self) -> QImage | None:
        """Render a clean, high-resolution sheet image for dependable printing.

        Qt's PDF/vector painter can replace Chinese glyphs with empty squares
        on some Windows systems. We therefore render the sheet viewport at five
        times its screen size: text, lines, and pasted pictures remain crisp,
        Chinese stays intact, and none of the editing-only UI is included.
        """

        if self.preview.rowCount() == 0:
            self.show_error("无法打印", "当前没有可打印的程序单预览。")
            return None
        table = self.preview
        old_size = table.size()
        old_minimum_size = table.minimumSize()
        old_maximum_size = table.maximumSize()
        old_horizontal = table.horizontalScrollBar().value()
        old_vertical = table.verticalScrollBar().value()
        old_horizontal_header = table.horizontalHeader().isVisible()
        old_vertical_header = table.verticalHeader().isVisible()
        old_horizontal_policy = table.horizontalScrollBarPolicy()
        old_vertical_policy = table.verticalScrollBarPolicy()
        selected_ranges = table.selectedRanges()
        current_row, current_column = table.currentRow(), table.currentColumn()
        selected_image = table.selected_image_id
        try:
            # The print sheet contains only Excel content: no A/B headers,
            # row numbers, scroll bars, selection blue or resize handles.
            table.clearSelection()
            table.clear_image_selection()
            table.horizontalHeader().setVisible(False)
            table.verticalHeader().setVisible(False)
            table.setHorizontalScrollBarPolicy(Qt.ScrollBarPolicy.ScrollBarAlwaysOff)
            table.setVerticalScrollBarPolicy(Qt.ScrollBarPolicy.ScrollBarAlwaysOff)
            table.horizontalScrollBar().setValue(0)
            table.verticalScrollBar().setValue(0)
            content_width = sum(table.columnWidth(column) for column in range(table.columnCount()))
            content_height = sum(table.rowHeight(row) for row in range(table.rowCount()))
            # ``table`` belongs to a splitter layout, which would otherwise
            # immediately resize it back to the on-screen pane size. Hold the
            # complete sheet dimensions only while it is being rendered.
            table.setFixedSize(content_width + 2, content_height + 2)
            table._reposition_image_overlays()
            QApplication.processEvents()
            source = table.viewport().rect()
            if source.width() <= 0 or source.height() <= 0:
                raise ValueError("程序单没有可打印的内容")
            # Five times the on-screen resolution gives clean fine grid lines
            # on ordinary office printers, while the cap avoids excessive RAM
            # use for a custom oversized template.
            render_scale = 5
            max_pixels = 60_000_000
            while render_scale > 2 and source.width() * source.height() * render_scale * render_scale > max_pixels:
                render_scale -= 1
            image = QImage(source.width() * render_scale, source.height() * render_scale, QImage.Format.Format_ARGB32)
            image.fill(QColor("#FFFFFF"))
            painter = QPainter(image)
            try:
                painter.scale(render_scale, render_scale)
                table.viewport().render(painter, QPoint(0, 0))
            finally:
                painter.end()
            top, left, bottom, right = self._saved_print_area()
            x = sum(table.columnWidth(column) for column in range(left))
            y = sum(table.rowHeight(row) for row in range(top))
            width = sum(table.columnWidth(column) for column in range(left, right + 1))
            height = sum(table.rowHeight(row) for row in range(top, bottom + 1))
            cropped = image.copy(x * render_scale, y * render_scale, width * render_scale, height * render_scale)
            # QTableWidget's grid does not always draw an outside edge. Add a
            # stable printed border like Excel/WPS around the chosen region.
            border_painter = QPainter(cropped)
            try:
                # Keep the whole stroke inside the image. Drawing a pen on
                # the zero edge cuts half of it off, which is especially
                # visible on a physical printer's left edge.
                border_width = max(4, render_scale + 1)
                inset = (border_width + 1) // 2
                border_painter.setPen(QPen(QColor("#000000"), border_width))
                border_painter.drawRect(
                    inset,
                    inset,
                    max(1, cropped.width() - inset * 2 - 1),
                    max(1, cropped.height() - inset * 2 - 1),
                )
            finally:
                border_painter.end()
            return cropped
        finally:
            table.setMinimumSize(old_minimum_size)
            table.setMaximumSize(old_maximum_size)
            table.resize(old_size)
            table.horizontalHeader().setVisible(old_horizontal_header)
            table.verticalHeader().setVisible(old_vertical_header)
            table.setHorizontalScrollBarPolicy(old_horizontal_policy)
            table.setVerticalScrollBarPolicy(old_vertical_policy)
            table.horizontalScrollBar().setValue(old_horizontal)
            table.verticalScrollBar().setValue(old_vertical)
            table._reposition_image_overlays()
            table.clearSelection()
            for selected_range in selected_ranges:
                table.setRangeSelected(selected_range, True)
            if current_row >= 0 and current_column >= 0:
                table.setCurrentCell(current_row, current_column)
            if selected_image:
                table.select_image(selected_image)

    @staticmethod
    def _paint_sheet_on_printer(printer: QPrinter, image: QImage, margin_mm: int = 8) -> None:
        """Fit the clean high-resolution sheet into the printable page centre."""

        painter = QPainter(printer)
        try:
            painter.setRenderHint(QPainter.RenderHint.SmoothPixmapTransform)
            # The printer painter uses device pixels. Supplying an explicit
            # target rectangle avoids QImage's own DPI metadata changing the
            # apparent size on PDF versus a physical Windows printer.
            page = printer.pageRect(QPrinter.Unit.DevicePixel)
            margin_pixels = max(0, min(30, margin_mm)) * printer.resolution() / 25.4
            # With the default ``fullPage=False``, Qt's painter origin is
            # already the top-left of ``pageRect``. Adding ``page.x/y`` here
            # applies a printer-driver margin a second time, visibly pushing
            # the sheet sideways and clipping its outside border.
            printable = QRectF(
                margin_pixels,
                margin_pixels,
                max(1.0, page.width() - margin_pixels * 2),
                max(1.0, page.height() - margin_pixels * 2),
            )
            scale = min(printable.width() / image.width(), printable.height() / image.height())
            width = image.width() * scale
            height = image.height() * scale
            target = QRectF(
                printable.x() + (printable.width() - width) / 2,
                printable.y() + (printable.height() - height) / 2,
                width,
                height,
            )
            painter.drawImage(target, image, QRectF(0, 0, image.width(), image.height()))
        finally:
            painter.end()

    def preview_program_sheet(self) -> None:
        """Show a wheel-zoomable print preview before printing."""

        image = self._capture_sheet_for_printing()
        if image is None:
            return
        printer = QPrinter(QPrinter.PrinterMode.HighResolution)
        printer.setPageSize(QPageSize(QPageSize.PageSizeId.A4))
        printer.setPageOrientation(self._configured_print_orientation())
        area_label = self._print_area_label(self._saved_print_area())
        dialog = ProgramSheetPrintPreviewDialog(
            printer,
            lambda target_printer: self._paint_sheet_on_printer(target_printer, image, self._configured_print_margin_mm()),
            area_label,
            self._configured_print_margin_mm(),
            self._save_print_preferences,
            self,
        )
        dialog.select_print_area_requested.connect(self.begin_print_area_selection)
        dialog.exec()

    def print_program_sheet(self) -> None:
        """Keep older callers working while always offering orientation preview."""

        self.preview_program_sheet()

    def _show_nc_import_menu(self, position: Any) -> None:
        """Provide a mouse alternative to Ctrl+V in the imported-files pane."""

        menu = QMenu(self.file_list)
        paste = menu.addAction("粘贴 NC 文件 / 程序（与 Ctrl+V 相同）")
        paste.triggered.connect(self._paste_nc_from_menu)
        menu.exec(self.file_list.mapToGlobal(position))

    def _paste_nc_from_menu(self) -> None:
        if not self.paste_nc_from_clipboard():
            self.statusBar().showMessage("剪贴板中没有可导入的 NC 文件或 NC 程序文本。", 4000)

    @staticmethod
    def _nc_paths_from_mime_data(mime_data: Any) -> list[Path]:
        """Return only supported local NC paths from a drag/drop payload."""

        if not mime_data.hasUrls():
            return []
        return [
            Path(url.toLocalFile())
            for url in mime_data.urls()
            if url.isLocalFile() and is_supported_nc_file(Path(url.toLocalFile()))
        ]

    def _clipboard_contains_nc_input(self) -> bool:
        mime_data = QApplication.clipboard().mimeData()
        return bool(clipboard_nc_paths(mime_data)) or (mime_data.hasText() and is_nc_program_text(mime_data.text()))

    def paste_nc_from_clipboard(self) -> bool:
        """Import Explorer-copied NC files or copied NC program text.

        Text is first saved in the per-user data directory, then parsed using
        the same code path as drag-and-drop. This keeps reparse available and
        avoids depending on clipboard contents after they are replaced.
        """

        mime_data = QApplication.clipboard().mimeData()
        paths = clipboard_nc_paths(mime_data)
        if paths:
            self.add_nc_files(paths)
            self.statusBar().showMessage("已从剪贴板导入 NC 文件。", 4000)
            return True
        text = mime_data.text() if mime_data.hasText() else ""
        if not is_nc_program_text(text):
            return False
        try:
            cached_program = cache_pasted_nc_program(text, clipboard_nc_filename(text))
            self.add_nc_files([cached_program])
            self.statusBar().showMessage("已从剪贴板导入 NC 程序文本。", 4000)
            return True
        except Exception as exc:
            self.show_error("无法粘贴 NC 程序", str(exc))
            return True

    def eventFilter(self, watched: Any, event: Any) -> bool:
        if watched in getattr(self, "_nc_drop_targets", ()) and event.type() in {
            QEvent.Type.DragEnter,
            QEvent.Type.DragMove,
            QEvent.Type.Drop,
        }:
            paths = self._nc_paths_from_mime_data(event.mimeData())
            if paths:
                if event.type() == QEvent.Type.Drop:
                    self.add_nc_files(paths)
                event.acceptProposedAction()
                return True
        if event.type() == QEvent.Type.KeyPress and event.matches(QKeySequence.StandardKey.Paste):
            if self._clipboard_contains_nc_input():
                self.paste_nc_from_clipboard()
                event.accept()
                return True
        return super().eventFilter(watched, event)

    def keyPressEvent(self, event: Any) -> None:
        if event.matches(QKeySequence.StandardKey.Paste) and self._clipboard_contains_nc_input():
            self.paste_nc_from_clipboard()
            event.accept()
            return
        super().keyPressEvent(event)

    def remove_images_at_selected_cell(self) -> None:
        selected_id = self.preview.selected_image_id
        if selected_id:
            self._remove_image_by_id(selected_id)
            return
        anchor = self.preview.current_coordinate()
        if not anchor:
            self.show_error("请选择图片位置", "请先选中图片，或在右侧 Excel 预览中点击图片所在的单元格。")
            return
        original_count = len(self.image_placements)
        self.image_placements = [placement for placement in self.image_placements if placement.anchor != anchor]
        if len(self.image_placements) == original_count:
            self.statusBar().showMessage(f"{anchor} 没有已插入的图片。", 3500)
            return
        self.preview.refresh_floating_images(self.image_placements)
        self.statusBar().showMessage(f"已移除 {anchor} 的图片。", 3500)

    def _remove_image_by_id(self, placement_id: str) -> None:
        original_count = len(self.image_placements)
        self.image_placements = [
            placement for placement in self.image_placements if placement.placement_id != placement_id
        ]
        if len(self.image_placements) == original_count:
            return
        self.preview.refresh_floating_images(self.image_placements)
        self.statusBar().showMessage("已移除选中的图片。", 3500)

    def choose_nc_files(self) -> None:
        filenames, _ = QFileDialog.getOpenFileNames(
            self,
            "选择 NC 程序",
            str(Path.home()),
            "NC 程序 (*.NC *.nc);;铜工无扩展名程序 (*)",
        )
        self.add_nc_files([Path(filename) for filename in filenames])

    def open_conversion_dialog(self, mode: TransformMode) -> None:
        """Open one legacy-shortcut-compatible conversion page.

        The dialog is intentionally the only place where a conversion writes
        to an NC source.  Once it returns successfully, the resulting file is
        parsed through the same conservative parser as a normally dropped
        program, so the existing program-sheet mapping remains unchanged.
        """

        dialog = NcConversionDialog(mode, self)
        if dialog.exec() != QDialog.DialogCode.Accepted:
            return
        self._import_converted_files(dialog.converted_files, mode)

    def _import_converted_files(self, converted_files: list[ConversionResult], mode: TransformMode) -> None:
        """Parse shortcut outputs and retain their pale-blue source labels.

        A transformed source may already be present in the program sheet.  In
        that case we replace only its existing rows in place rather than
        shifting the operator's later rows, inserted blanks, images or manual
        edits.  For M06, the reviewed pre-conversion snapshot retains the
        original tool-change boundaries for multi-tool program-sheet rows,
        while the disk file itself remains the exact shortcut output.
        """

        parsed_by_target: list[tuple[Path, Path, list[ProgramRecord], str | None]] = []
        for conversion in converted_files:
            source = conversion.source_path.resolve()
            target = conversion.target_path.resolve()
            try:
                if mode is TransformMode.M06 and conversion.parser_snapshot is not None:
                    parsed = apply_tool_number_display_policy(
                        parse_records_text(conversion.parser_snapshot, target.name, target)
                    )
                else:
                    parsed = apply_tool_number_display_policy(parse_file_records(target))
            except Exception as exc:
                self.show_error("NC 文件解析失败", f"{target.name}\n\n{exc}")
                continue
            parsed_by_target.append((source, target, parsed, conversion.parser_snapshot))

        if not parsed_by_target:
            return

        overflow_records: list[ProgramRecord] = []
        added_count = 0
        for source, target, parsed_records, parser_snapshot in parsed_by_target:
            # The converted output is the only source that should retain the
            # mode label. Copper conversion has a different, extensionless
            # output path, so remove the now-deleted .NC path from the map.
            self.conversion_tags.pop(source, None)
            self.conversion_tags[target] = CONVERSION_TAGS[mode]
            if mode is TransformMode.M06 and parser_snapshot is not None:
                self.m06_parse_snapshots[target] = parser_snapshot
            else:
                # A later M304/copper operation means the current disk source
                # is now authoritative, so an older M06 split snapshot must
                # not be reused.
                self.m06_parse_snapshots.pop(target, None)
            if source != target:
                self.m06_parse_snapshots.pop(source, None)

            existing_indices = [
                index
                for index, record in enumerate(self.records)
                if record is not None and record.source_path.resolve() in {source, target}
            ]
            if existing_indices:
                for position, index in enumerate(existing_indices):
                    self.records[index] = parsed_records[position] if position < len(parsed_records) else None
                    if position >= len(parsed_records):
                        # An old tool section disappeared during conversion;
                        # preserve the row position as a deliberate blank.
                        self.row_markers.pop(index, None)

                extra_records = parsed_records[len(existing_indices) :]
                free_rows = LAYOUT.capacity - len(self.records)
                if extra_records and free_rows > 0:
                    self.records.extend(extra_records[:free_rows])
                    added_count += min(len(extra_records), free_rows)
                overflow_records.extend(extra_records[free_rows:])
                continue

            # This source was not already in the sheet. Append its parsed
            # rows in the same way as an ordinary drop, retaining the dialog's
            # natural filename ordering.
            free_rows = LAYOUT.capacity - len(self.records)
            if free_rows > 0:
                self.records.extend(parsed_records[:free_rows])
                added_count += min(len(parsed_records), free_rows)
            overflow_records.extend(parsed_records[free_rows:])

        if overflow_records:
            self.show_capacity_overflow(overflow_records)
        self._sort_records_by_program_name_if_safe()
        self._apply_records_to_view()
        converted_names = "、".join(target.name for _, target, _, _ in parsed_by_target[:4])
        remaining = len(parsed_by_target) - 4
        if remaining > 0:
            converted_names += f" 等 {len(parsed_by_target)} 个文件"
        if added_count:
            self.statusBar().showMessage(f"已按原 {CONVERSION_TAGS[mode]} 快捷方式修改并导入：{converted_names}", 5000)
        else:
            self.statusBar().showMessage(f"已按原 {CONVERSION_TAGS[mode]} 快捷方式修改并重新解析：{converted_names}", 5000)

    def edit_imported_nc_file(self, item: QListWidgetItem) -> None:
        """Let an operator inspect and edit the source behind a list entry."""

        source_value = item.data(Qt.ItemDataRole.UserRole)
        if not source_value:
            return
        source_path = Path(str(source_value))
        if not source_path.is_file():
            self.show_error(
                "找不到 NC 程序",
                f"{source_path}\n\n该文件可能已被外部移动或删除。请重新导入正确的 NC 文件。",
            )
            return
        try:
            dialog = NcSourceEditorDialog(source_path, self)
        except Exception as exc:
            self.show_error("无法打开 NC 程序", f"{source_path.name}\n\n{exc}")
            return
        if dialog.exec() == QDialog.DialogCode.Accepted:
            if dialog.did_change:
                # A changed M06 output no longer has reliably recoverable
                # original section boundaries, so parse the edited source as
                # it now exists on disk rather than ignoring the edit.
                self.m06_parse_snapshots.pop(source_path.resolve(), None)
            self.reparse()
            self.statusBar().showMessage(f"已保存并重新解析：{source_path.name}", 4500)

    def dragEnterEvent(self, event: Any) -> None:  # Qt event type differs across bindings
        if self._nc_paths_from_mime_data(event.mimeData()):
            event.acceptProposedAction()

    def dragMoveEvent(self, event: Any) -> None:
        if self._nc_paths_from_mime_data(event.mimeData()):
            event.acceptProposedAction()

    def dropEvent(self, event: Any) -> None:
        paths = self._nc_paths_from_mime_data(event.mimeData())
        if paths:
            self.add_nc_files(paths)
            event.acceptProposedAction()

    @staticmethod
    def _program_row_index(table_row: int) -> int | None:
        index = table_row - (LAYOUT.first_program_row - 1)
        return index if 0 <= index < LAYOUT.capacity else None

    def _has_imported_records(self) -> bool:
        return any(record is not None for record in self.records)

    def _sort_records_by_program_name_if_safe(self) -> bool:
        """Keep ordinary imports in natural filename order without moving edits.

        The first import and a plain conversion can safely be re-ordered.  Once
        an operator has inserted a spacer/marker, added an image, or typed a
        value that is not represented by its ``ProgramRecord``, preserving the
        existing physical rows is safer than silently moving their work.
        """

        if len(self.records) < 2 or any(record is None for record in self.records):
            return False
        if self.row_markers or self.image_placements:
            return False
        automatic = automatic_cells_for_rows(self.records)
        for coordinate, value in self.manual_cells.items():
            row_text = "".join(character for character in coordinate if character.isdigit())
            if not row_text:
                continue
            row = int(row_text)
            if not LAYOUT.first_program_row <= row <= LAYOUT.last_program_row:
                continue
            expected = automatic.get(coordinate)
            if expected is None:
                if value:
                    return False
            # A blank is also how the standard program-area reset represents
            # a newly added row before `_apply_records_to_view` populates it.
            # Non-blank differences are the meaningful operator edits that
            # must keep their current row position.
            elif value and value != expected:
                return False

        sorted_records = sorted(
            (record for record in self.records if record is not None),
            key=lambda record: (natural_program_sort_key(record.program_name), record.source_tool_index),
        )
        if list(self.records) == sorted_records:
            return False
        self.records = sorted_records
        return True

    def _has_tracked_data(self, index: int) -> bool:
        """Whether moving this program row would discard a user-entered value."""

        excel_row = LAYOUT.first_program_row + index
        for column in range(2, self.preview.columnCount() + 1):
            coordinate = f"{ExcelPreview._column_name(column)}{excel_row}"
            if self.manual_cells.get(coordinate, ""):
                return True
        return False

    def _has_image_at_program_index(self, index: int) -> bool:
        row = LAYOUT.first_program_row + index
        return any(
            "".join(character for character in placement.anchor if character.isdigit()) == str(row)
            for placement in self.image_placements
        )

    def _has_only_marker_data(self, index: int) -> bool:
        """Whether a marker row has no extra operator-entered cells."""

        marker = self.row_markers.get(index)
        if marker is None:
            return False
        marker_cell = LAYOUT.target_cell("program_name", index)
        excel_row = LAYOUT.first_program_row + index
        for column in range(2, self.preview.columnCount() + 1):
            coordinate = f"{ExcelPreview._column_name(column)}{excel_row}"
            value = self.manual_cells.get(coordinate, "")
            if coordinate == marker_cell and value == marker:
                continue
            if value:
                return False
        return True

    def _shift_program_images(self, start_index: int, direction: int) -> None:
        """Keep image anchors aligned when a logical program row moves."""

        shifted: list[ImagePlacement] = []
        for placement in self.image_placements:
            letters = "".join(character for character in placement.anchor if character.isalpha())
            digits = "".join(character for character in placement.anchor if character.isdigit())
            if not digits:
                shifted.append(placement)
                continue
            index = int(digits) - LAYOUT.first_program_row
            if 0 <= index < LAYOUT.capacity and index >= start_index:
                destination = index + direction
                if 0 <= destination < LAYOUT.capacity:
                    shifted.append(
                        replace(
                            placement,
                            anchor=f"{letters}{LAYOUT.first_program_row + destination}",
                        )
                    )
            else:
                shifted.append(placement)
        self.image_placements = shifted
        self.preview.refresh_floating_images(self.image_placements)

    def _shift_tracked_program_cells(self, start_index: int, direction: int) -> None:
        """Shift exported manual values in lockstep with the preview values."""

        moved: dict[str, str] = {}
        remove: list[str] = []
        first = LAYOUT.first_program_row
        last = LAYOUT.last_program_row
        for coordinate, value in self.manual_cells.items():
            letters = "".join(character for character in coordinate if character.isalpha())
            digits = "".join(character for character in coordinate if character.isdigit())
            if not digits:
                continue
            row = int(digits)
            index = row - first
            if first <= row <= last and index >= start_index:
                destination = index + direction
                if 0 <= destination < LAYOUT.capacity:
                    moved[f"{letters}{first + destination}"] = value
                remove.append(coordinate)
        for coordinate in remove:
            self.manual_cells.pop(coordinate, None)
        self.manual_cells.update(moved)

    def _shift_marker_rows_for_insert(self, index: int) -> None:
        self.row_markers = {
            row_index + 1 if row_index >= index else row_index: marker
            for row_index, marker in self.row_markers.items()
            if row_index + 1 < LAYOUT.capacity or row_index < index
        }

    def _shift_marker_rows_for_remove(self, index: int) -> None:
        self.row_markers = {
            row_index - 1 if row_index > index else row_index: marker
            for row_index, marker in self.row_markers.items()
            if row_index != index
        }

    def _insert_program_row(self, table_row: int, marker: str | None = None) -> None:
        """Insert either a blank row or a red machining-orientation marker."""

        index = self._program_row_index(table_row)
        if index is None:
            return
        if not self._has_imported_records():
            self.statusBar().showMessage("请先导入 NC 程序，再插入程序单行。", 3500)
            return
        if self._has_tracked_data(LAYOUT.capacity - 1) or self._has_image_at_program_index(LAYOUT.capacity - 1) or (
            len(self.records) == LAYOUT.capacity and self.records[-1] is not None
        ):
            self.show_error("无法插入程序行", "模板最后一行已有数据。请先腾出最后一个程序行，避免覆盖数据。")
            return
        # A row below the final imported program can still be selected in the
        # preview. Preserve its requested position by representing preceding
        # gaps explicitly.
        while len(self.records) < index:
            self.records.append(None)
        if len(self.records) == LAYOUT.capacity:
            self.records.pop()
            self.row_markers.pop(LAYOUT.capacity - 1, None)
        self.preview.shift_program_row_values(table_row, 1)
        self._shift_tracked_program_cells(index, 1)
        self._shift_program_images(index, 1)
        self._shift_marker_rows_for_insert(index)
        # ``shift_program_row_values`` has moved the old NC-generated values
        # down with the rest of the visible row.  A newly inserted spacer must
        # explicitly clear only those generated columns; other operator edits
        # elsewhere in the sheet remain untouched on later re-parses.
        for field in LAYOUT.columns:
            coordinate = LAYOUT.target_cell(field, index)
            self.manual_cells[coordinate] = ""
            self.preview.set_value(coordinate, "")
        self.records.insert(index, None)
        if marker:
            self.row_markers[index] = marker
        self._apply_records_to_view()
        if marker:
            self.statusBar().showMessage(f"已在所选行下方插入红色提示行：{marker}。选中该行后按 Delete 可取消。", 4500)
        else:
            self.statusBar().showMessage("已在所选行下方插入空行；选中该空行后按 Delete 可取消。", 4500)

    def insert_blank_program_row(self, table_row: int) -> None:
        """Insert an empty logical row below the selected preview row."""

        self._insert_program_row(table_row)

    def insert_marker_program_row(self, table_row: int, marker: str) -> None:
        if marker not in ROTATION_MARKERS:
            self.show_error("无法插入提示行", "提示行内容无效。")
            return
        self._insert_program_row(table_row, marker)

    def remove_blank_program_row(self, table_row: int) -> None:
        """Remove an operator-inserted empty row and close the gap."""

        index = self._program_row_index(table_row)
        if index is None or index >= len(self.records) or self.records[index] is not None:
            self.statusBar().showMessage("请选择通过“插入空行”创建的空白程序行，再按 Delete。", 4000)
            return
        has_extra_data = self._has_tracked_data(index) and not self._has_only_marker_data(index)
        if has_extra_data or self._has_image_at_program_index(index):
            self.statusBar().showMessage("该行已有手动填写内容或图片，无法作为空行取消。请先清空该行内容。", 4000)
            return
        self.preview.shift_program_row_values(table_row, -1)
        self._shift_tracked_program_cells(index, -1)
        self._shift_program_images(index, -1)
        self._shift_marker_rows_for_remove(index)
        self.records.pop(index)
        self._apply_records_to_view()
        self.statusBar().showMessage("已取消空行或提示行，后续程序数据已上移。", 3500)

    def delete_preview_cells(self, top: int, left: int, bottom: int, right: int) -> None:
        """Keep Delete compatible with both blank-row removal and Excel edits."""

        index = self._program_row_index(top)
        if top == bottom and index is not None and index < len(self.records) and self.records[index] is None:
            self.remove_blank_program_row(top)
            return
        self.preview.clear_selected_cells()

    def add_nc_files(self, paths: list[Path]) -> None:
        paths = [path for path in paths if is_supported_nc_file(path)]
        if not paths:
            self.show_error("未导入文件", "请拖入或选择至少一个 .NC / .nc 文件，或无扩展名的铜工 NC 程序。")
            return
        paths.sort(key=lambda path: natural_program_sort_key(path.name))
        existing = {record.source_path.resolve() for record in self.records if record is not None}
        overflow_records: list[ProgramRecord] = []
        for path in paths:
            if path.resolve() in existing:
                continue
            try:
                parsed_records = apply_tool_number_display_policy(parse_file_records(path))
                free_rows = LAYOUT.capacity - len(self.records)
                if free_rows <= 0:
                    overflow_records.extend(parsed_records)
                    continue
                self.records.extend(parsed_records[:free_rows])
                overflow_records.extend(parsed_records[free_rows:])
                existing.add(path.resolve())
            except Exception as exc:
                self.show_error("NC 文件解析失败", f"{path.name}\n\n{exc}")
        if overflow_records:
            self.show_capacity_overflow(overflow_records)
        self._sort_records_by_program_name_if_safe()
        self._apply_records_to_view()

    def show_capacity_overflow(self, overflow_records: list[ProgramRecord]) -> None:
        """Name every omitted program row and give a safe continuation path."""

        labels = [f"{record.program_name}{f'（{record.tool_number}）' if record.tool_number else ''}" for record in overflow_records]
        preview_labels = labels[:20]
        displayed = "\n".join(f"• {label}" for label in preview_labels)
        if len(labels) > len(preview_labels):
            displayed += f"\n• ……其余 {len(labels) - len(preview_labels)} 个请点“显示详细信息”查看"
        box = QMessageBox(self)
        box.setWindowTitle("模板容量已满")
        box.setIcon(QMessageBox.Icon.Warning)
        box.setText(f"模板最多填写 {LAYOUT.capacity} 个程序行。以下 {len(labels)} 个程序行本次未导入：")
        box.setInformativeText(
            f"{displayed}\n\n"
            "请先点击“完成并保存”保存当前已整理的程序单，然后点击“清空”清除当前导入记录，"
            "再重新导入上述未导入的程序。\n\n"
            "“清空”只清除软件当前列表，不会删除任何原始 NC 文件。"
        )
        box.setDetailedText("未导入的程序行：\n" + "\n".join(labels))
        box.exec()

    def reparse(self) -> None:
        if not self._has_imported_records():
            self.statusBar().showMessage("当前没有可重新解析的 NC 文件。", 3500)
            return
        refreshed_by_path: dict[Path, list[ProgramRecord]] = {}
        for record in self.records:
            if record is None:
                continue
            source = record.source_path.resolve()
            if source not in refreshed_by_path:
                try:
                    snapshot = self.m06_parse_snapshots.get(source)
                    if snapshot is not None:
                        refreshed_by_path[source] = apply_tool_number_display_policy(
                            parse_records_text(snapshot, record.source_path.name, record.source_path)
                        )
                    else:
                        refreshed_by_path[source] = apply_tool_number_display_policy(parse_file_records(record.source_path))
                except Exception as exc:
                    refreshed_by_path[source] = [
                        ProgramRecord(
                            record.source_path,
                            record.program_name,
                            source_tool_index=record.source_tool_index,
                            warnings=[f"重新解析失败：{exc}"],
                        )
                    ]
        refreshed: list[ProgramRecord | None] = []
        for record in self.records:
            if record is None:
                refreshed.append(None)
                continue
            parsed_records = refreshed_by_path[record.source_path.resolve()]
            if record.source_tool_index < len(parsed_records):
                refreshed.append(parsed_records[record.source_tool_index])
            else:
                refreshed.append(
                    ProgramRecord(
                        record.source_path,
                        record.program_name,
                        source_tool_index=record.source_tool_index,
                        warnings=["重新解析后未找到对应刀具段"],
                    )
                )
        apply_tool_number_display_policy([record for record in refreshed if record is not None])
        self.records = refreshed
        self._apply_records_to_view()
        self.statusBar().showMessage("已按原始 NC 文件重新解析。", 3500)

    def clear_all(self) -> None:
        if (
            self._has_imported_records()
            and self.settings.get("confirm_before_clear", True)
            and QMessageBox.question(self, "清空", "清空已导入文件及当前预览中的手动修改？") != QMessageBox.StandardButton.Yes
        ):
            return
        self.records.clear()
        self.row_markers.clear()
        self.conversion_tags.clear()
        self.m06_parse_snapshots.clear()
        self.manual_cells.clear()
        self.image_placements.clear()
        self._program_area_reset = False
        self.file_list.clear()
        self._set_result_table()
        if self.template_path:
            self.preview.load_template(self.template_path)
            self.preview.refresh_floating_images(self.image_placements)
            self.manual_cells.update(automatic_cells_for_rows([]))
            self.preview.set_value(LAYOUT.date_cell, self.manual_cells[LAYOUT.date_cell])
        self.statusBar().showMessage("已恢复到未填写的模板预览。", 3500)

    def _apply_records_to_view(self) -> None:
        self.file_list.clear()
        listed_sources: set[Path] = set()
        for record in self.records:
            if record is None:
                continue
            source = record.source_path.resolve()
            if source in listed_sources:
                continue
            item = QListWidgetItem(record.program_name)
            item.setData(Qt.ItemDataRole.UserRole, str(source))
            item.setToolTip(f"双击查看和手动修改 NC 程序：\n{source}")
            conversion_tag = self.conversion_tags.get(source)
            if conversion_tag:
                item.setText(f"{conversion_tag}    {record.program_name}")
                item.setBackground(CONVERSION_TAG_BACKGROUND)
                item.setForeground(CONVERSION_TAG_FOREGROUND)
                item_font = item.font()
                item_font.setBold(True)
                item.setFont(item_font)
            self.file_list.addItem(item)
            listed_sources.add(source)
        # The supplied workbook contains filled reference rows. On the first
        # import, blank only the NC-generated fields across the program area so
        # old sample entries cannot leak into the exported sheet.
        if not self._program_area_reset:
            cleared = cleared_program_cells()
            self.manual_cells.update(cleared)
            for cell, value in cleared.items():
                self.preview.set_value(cell, value)
            self._program_area_reset = True
        # Re-parsing deliberately refreshes values that come from the NC
        # records. Do *not* blank the entire program area here: the operator
        # may have typed remarks, copied notes or pasted a range into other
        # editable cells, and those values must remain visible after any edit,
        # re-parse or additional NC import.
        automatic = automatic_cells_for_rows(self.records)
        self.manual_cells.update(automatic)
        for cell, value in automatic.items():
            self.preview.set_value(cell, value)
        # Normal data returns to the template's black text. Marker cells are
        # reapplied below in red after automatic values have cleared the area.
        for index in range(LAYOUT.capacity):
            self.preview.set_text_color(LAYOUT.target_cell("program_name", index), QColor("#000000"))
        for index, marker in self.row_markers.items():
            coordinate = LAYOUT.target_cell("program_name", index)
            self.manual_cells[coordinate] = marker
            self.preview.set_value(coordinate, marker)
            self.preview.set_text_color(coordinate, MARKER_RED)
        # Imported/reparsed NC values are a new baseline, so undo only ever
        # affects the operator's edits made after this refresh.
        self.preview.clear_edit_history()
        self._set_result_table()

    def _set_result_table(self) -> None:
        self._result_loading = True
        self.result_table.setRowCount(len(self.records))
        for row, record in enumerate(self.records):
            for column, (_, field) in enumerate(RESULT_COLUMNS):
                if record is None:
                    marker = self.row_markers.get(row)
                    if field == "program_name" and marker:
                        value = marker
                    elif field == "status":
                        value = "红色提示行（预览中按 Delete 取消）" if marker else "空行（预览中按 Delete 取消）"
                    else:
                        value = ""
                else:
                    value = record.status if field == "status" else str(getattr(record, field))
                item = QTableWidgetItem(value)
                if field == "status" or record is None:
                    item.setFlags(item.flags() & ~Qt.ItemFlag.ItemIsEditable)
                    if record is not None and record.has_unrecognized_values:
                        item.setForeground(QColor("#a15c00"))
                    elif record is None and self.row_markers.get(row):
                        item.setForeground(MARKER_RED)
                self.result_table.setItem(row, column, item)
        self.result_table.resizeColumnsToContents()
        self._result_loading = False

    def _on_result_changed(self, item: QTableWidgetItem) -> None:
        if self._result_loading or item.row() >= len(self.records):
            return
        field = RESULT_COLUMNS[item.column()][1]
        if field not in FIELD_NAMES:
            return
        record = self.records[item.row()]
        if record is None:
            return
        setattr(record, field, item.text())
        target = LAYOUT.target_cell(field, item.row())
        self.manual_cells[target] = item.text()
        self.preview.set_value(target, item.text())

    def _on_preview_edited(self, cell: str, value: str) -> None:
        self.manual_cells[cell] = value
        if cell.startswith("B"):
            row_text = "".join(character for character in cell if character.isdigit())
            index = int(row_text) - LAYOUT.first_program_row if row_text else -1
            if index in self.row_markers:
                if value:
                    self.row_markers[index] = value
                    self.preview.set_text_color(cell, MARKER_RED)
                else:
                    self.row_markers.pop(index, None)
                    self.preview.set_text_color(cell, QColor("#000000"))
                self._set_result_table()
                return
        for index, record in enumerate(self.records):
            if record is None:
                continue
            for field in FIELD_NAMES:
                if LAYOUT.target_cell(field, index) == cell:
                    setattr(record, field, value)
                    self._set_result_table()
                    return

    def _repository(self) -> str:
        return str(self.settings.get("update_repository") or load_embedded_update_repository()).strip()

    def _apply_toolbar_button_visibility(self) -> None:
        """Apply the user's compact-toolbar choices without losing actions."""

        saved = self.settings.get("toolbar_button_visibility", {})
        visibility = saved if isinstance(saved, dict) else {}
        for key, button in self._optional_toolbar_buttons.items():
            button.setVisible(bool(visibility.get(key, True)))

    def edit_settings(self) -> None:
        """Edit ordinary user preferences from the single Settings button."""

        dialog = SettingsDialog(
            self._repository(),
            bool(self.settings.get("show_save_success_message", True)),
            bool(self.settings.get("automatic_update_check", True)),
            bool(self.settings.get("confirm_before_clear", True)),
            bool(self.settings.get("confirm_before_save", False)),
            bool(self.settings.get("open_folder_after_save", False)),
            self.settings.get("toolbar_button_visibility", {}) if isinstance(self.settings.get("toolbar_button_visibility", {}), dict) else {},
            str(self.settings.get("print_orientation", "portrait")),
            self._configured_print_margin_mm(),
            str(self.settings.get("selection_color", DEFAULT_SELECTION_COLOR)),
            self,
        )
        if dialog.exec() == QDialog.DialogCode.Accepted:
            self.settings["update_repository"] = dialog.repository.text().strip()
            previous_reminder = bool(self.settings.get("show_save_success_message", True))
            self.settings["show_save_success_message"] = dialog.save_reminder.isChecked()
            self.settings["automatic_update_check"] = dialog.automatic_update_check.isChecked()
            self.settings["confirm_before_clear"] = dialog.confirm_before_clear.isChecked()
            self.settings["confirm_before_save"] = dialog.confirm_before_save.isChecked()
            self.settings["open_folder_after_save"] = dialog.open_folder_after_save.isChecked()
            self.settings["toolbar_button_visibility"] = {
                key: checkbox.isChecked() for key, checkbox in dialog.toolbar_buttons.items()
            }
            self.settings["print_orientation"] = str(dialog.print_orientation.currentData())
            self.settings["print_margin_mm"] = dialog.print_margin_mm.value()
            self.settings["selection_color"] = dialog.selected_selection_color()
            # Re-enabling save notifications should take effect on the next
            # save even if this build has already shown its first-save notice.
            if dialog.save_reminder.isChecked() and not previous_reminder:
                self.settings.pop("save_notice_version", None)
            save_settings(self.settings)
            self._apply_toolbar_button_visibility()
            self.preview.set_selection_color(self.settings["selection_color"])
            self.statusBar().showMessage("设置已保存。", 3500)

    def open_find_replace(self, replace_mode: bool = True) -> None:
        """Show non-destructive Excel-style find/replace for the sheet preview."""

        dialog = FindReplaceDialog(self.preview, replace_mode, self)
        dialog.exec()

    # Keep the old internal name available for any older integrations.
    def edit_update_settings(self) -> None:
        self.edit_settings()

    def _automatic_update_check(self) -> None:
        if self.settings.get("automatic_update_check", True) and self._repository():
            self.check_updates(silent=True)

    def check_updates(self, silent: bool) -> None:
        if self._update_worker and self._update_worker.isRunning():
            return
        repository = self._repository()
        if not repository:
            if not silent:
                self.show_error("尚未配置更新源", "请先点击“更新设置”，填写 GitHub 仓库 owner/repository。")
            return
        self.statusBar().showMessage("正在检查 GitHub Release 更新…")
        worker = Worker(lambda: check_latest(repository), self)
        self._update_worker = worker
        worker.succeeded.connect(lambda update: self._on_update_check_success(update, silent))
        worker.failed.connect(lambda error: self._on_update_check_failed(error, silent))
        worker.finished.connect(worker.deleteLater)
        worker.start()

    def _on_update_check_success(self, update: UpdateInfo | None, silent: bool) -> None:
        self._update_worker = None
        if update is None:
            if not silent:
                self.statusBar().showMessage("已是最新版本。", 3500)
            return
        box = QMessageBox(self)
        box.setWindowTitle("发现新版本")
        box.setIcon(QMessageBox.Icon.Information)
        box.setText(f"发现 v{update.version}，当前为 v{__version__}。")
        box.setInformativeText("更新只替换程序安装文件；用户模板、配置和个人数据均保存在 AppData，不会被覆盖。")
        install = box.addButton("下载并安装", QMessageBox.ButtonRole.AcceptRole)
        release = box.addButton("查看 Release", QMessageBox.ButtonRole.ActionRole)
        box.addButton("稍后", QMessageBox.ButtonRole.RejectRole)
        box.exec()
        if box.clickedButton() is release:
            open_release_page(update)
        elif box.clickedButton() is install:
            self._download_update(update)

    def _on_update_check_failed(self, error: str, silent: bool) -> None:
        self._update_worker = None
        if not silent:
            self.show_error("检查更新失败", error)

    def _download_update(self, update: UpdateInfo) -> None:
        self.statusBar().showMessage("正在下载并校验更新安装包…")
        worker = Worker(lambda: download_verified_installer(update), self)
        self._update_worker = worker
        worker.succeeded.connect(self._on_update_downloaded)
        worker.failed.connect(lambda error: self._on_update_check_failed(error, False))
        worker.finished.connect(worker.deleteLater)
        worker.start()

    def _on_update_downloaded(self, installer: Path) -> None:
        self._update_worker = None
        try:
            start_installer_after_exit(installer)
            QMessageBox.information(self, "开始升级", "已下载并校验更新。程序关闭后将自动安装新版。")
            self._exiting_for_update = True
            QApplication.instance().quit()
        except Exception as exc:
            self.show_error("无法启动更新安装程序", str(exc))

    @staticmethod
    def _default_save_directory() -> Path:
        """Return Desktop only as the initial suggestion, never as a forced path."""

        from PySide6.QtCore import QStandardPaths

        desktop_text = QStandardPaths.writableLocation(QStandardPaths.StandardLocation.DesktopLocation)
        if desktop_text:
            return Path(desktop_text)
        return Path(os.path.expanduser("~")) / "Desktop"

    def _configured_save_directory(self) -> Path | None:
        configured = str(self.settings.get("output_directory", "")).strip()
        if not configured:
            return None
        directory = Path(configured).expanduser()
        return directory if directory.is_dir() else None

    def configure_save_directory(self) -> Path | None:
        """Ask once for a persistent, user-owned generated-sheet directory."""

        current_directory = self._configured_save_directory() or self._default_save_directory()
        dialog = SaveDirectoryDialog(current_directory, self)
        if dialog.exec() != QDialog.DialogCode.Accepted:
            return None
        raw_directory = dialog.directory.text().strip()
        if not raw_directory:
            self.show_error("未设置保存位置", "请输入或选择一个程序单保存文件夹。")
            return None
        directory = Path(raw_directory).expanduser()
        try:
            directory.mkdir(parents=True, exist_ok=True)
        except OSError as exc:
            self.show_error("无法创建保存文件夹", str(exc))
            return None
        if not directory.is_dir():
            self.show_error("无法设置保存位置", f"该位置不是文件夹：{directory}")
            return None
        self.settings["output_directory"] = str(directory)
        save_settings(self.settings)
        self.statusBar().showMessage(f"已设置固定保存位置：{directory}", 5000)
        return directory

    def _validate_before_save(self) -> bool:
        if not self._has_imported_records():
            self.show_error("无法保存", "请先拖入或选择至少一个 NC 程序。")
            return False
        try:
            automatic_cells_for_rows(self.records)
        except TemplateError as exc:
            self.show_error("无法保存", str(exc))
            return False
        uncertain = [record.program_name for record in self.records if record is not None and record.has_unrecognized_values]
        if uncertain:
            text = "以下程序仍有“未识别”项：\n" + "、".join(uncertain)
            text += "\n\n可在左侧解析表或右侧 Excel 预览中手动修改，仍要保存吗？"
            return QMessageBox.question(self, "需要确认", text) == QMessageBox.StandardButton.Yes
        return True

    def save_output(self) -> None:
        if not self.template_path:
            return
        # On a fresh installation there is no saved folder, so the first
        # press of “完成并保存” prompts for it exactly once.
        output_directory = self._configured_save_directory() or self.configure_save_directory()
        if output_directory is None or not self._validate_before_save():
            return
        try:
            # Ensure table edits are applied even when a preview cell was never clicked.
            self.manual_cells.update(automatic_cells_for_rows(self.records))
            target = output_filename(
                [record for record in self.records if record is not None],
                output_directory,
                document_name=self.manual_cells.get(LAYOUT.output_name_cell, ""),
            )
            if self.settings.get("confirm_before_save", False):
                answer = QMessageBox.question(self, "确认保存", f"确认保存程序单到：\n{target}？")
                if answer != QMessageBox.StandardButton.Yes:
                    return
            export_workbook(
                self.template_path,
                target,
                self.manual_cells,
                images=self.image_placements,
                red_text_cells=[LAYOUT.target_cell("program_name", index) for index in self.row_markers],
            )
            self._show_save_success(target)
            if self.settings.get("open_folder_after_save", False) and sys.platform == "win32":
                os.startfile(str(target.parent))
        except Exception as exc:
            self.show_error("保存程序单失败", str(exc))

    def _show_save_success(self, target: Path) -> None:
        # A newly installed version explains the completed first save once.
        # Later saves stay quiet (status bar only), and “永不提醒” suppresses
        # this message even after future upgrades.
        if self.settings.get("show_save_success_message", True) is False:
            self.statusBar().showMessage(f"程序单已保存：{target}", 5000)
            return
        if self.settings.get("save_notice_version") == __version__:
            self.statusBar().showMessage(f"程序单已保存：{target}", 5000)
            return
        self.settings["save_notice_version"] = __version__
        save_settings(self.settings)
        box = QMessageBox(self)
        box.setWindowTitle("新版首次保存完成")
        box.setIcon(QMessageBox.Icon.Information)
        box.setText(f"v{__version__} 已完成第一次程序单保存")
        box.setInformativeText(f"{target}\n\n之后保存将不再弹窗，只会在底部显示完成信息。")
        never_remind = QCheckBox("永不提醒")
        box.setCheckBox(never_remind)
        open_file = box.addButton("打开文件", QMessageBox.ButtonRole.AcceptRole)
        open_folder = box.addButton("打开文件夹", QMessageBox.ButtonRole.ActionRole)
        box.addButton("关闭", QMessageBox.ButtonRole.RejectRole)
        box.exec()
        if never_remind.isChecked():
            self.settings["show_save_success_message"] = False
            save_settings(self.settings)
        if box.clickedButton() is open_file:
            os.startfile(target)  # type: ignore[attr-defined]
        elif box.clickedButton() is open_folder:
            os.startfile(target.parent)  # type: ignore[attr-defined]

    def closeEvent(self, event: Any) -> None:
        """Ask a visible user whether X should minimize or quit.

        The remembered choice intentionally lives with other user settings in
        LocalAppData, so upgrades never reset it.  Unit-test windows are not
        shown; accepting those invisible closes keeps test cleanup non-modal.
        """

        application = QApplication.instance()
        # The offscreen Qt backend is used only by automated tests. It has no
        # taskbar to minimize to and must never wait for a human response.
        if (
            self._exiting_for_update
            or not self.isVisible()
            or (application is not None and application.platformName().casefold() == "offscreen")
        ):
            event.accept()
            return
        remembered = str(self.settings.get("close_window_behavior", ""))
        if remembered == "minimize":
            event.ignore()
            self._hide_to_tray()
            return
        if remembered == "exit":
            event.accept()
            return

        box = QMessageBox(self)
        box.setWindowTitle("关闭 CNC 智能程序单")
        box.setIcon(QMessageBox.Icon.Question)
        box.setText("要缩小到电脑右下角隐藏图标区，还是直接退出软件？")
        never_remind = QCheckBox("永不提醒")
        box.setCheckBox(never_remind)
        minimize = box.addButton("缩小到右下角隐藏图标区", QMessageBox.ButtonRole.AcceptRole)
        exit_program = box.addButton("直接退出", QMessageBox.ButtonRole.DestructiveRole)
        box.exec()
        selected = box.clickedButton()
        if selected is minimize:
            if never_remind.isChecked():
                self.settings["close_window_behavior"] = "minimize"
                save_settings(self.settings)
            event.ignore()
            self._hide_to_tray()
        elif selected is exit_program:
            if never_remind.isChecked():
                self.settings["close_window_behavior"] = "exit"
                save_settings(self.settings)
            event.accept()
        else:
            event.ignore()

    def show_error(self, title: str, details: str) -> None:
        log_path = log_exception(f"{title}\n{details}")
        QMessageBox.critical(self, title, f"{details}\n\n详细日志：{log_path}")


def run() -> int:
    _set_windows_app_user_model_id()
    app = QApplication(sys.argv)
    apply_application_theme(app)
    app.setApplicationName(APP_NAME)
    app.setOrganizationName("CNCProgramSheet")
    icon = QIcon(str(resource_path("assets", APP_ICON_FILE)))
    if not icon.isNull():
        app.setWindowIcon(icon)
    window = MainWindow()
    window.show()
    return app.exec()


def _exception_hook(exc_type: type[BaseException], value: BaseException, trace: Any) -> None:
    details = "".join(traceback.format_exception(exc_type, value, trace))
    log_path = log_exception(details)
    if QApplication.instance() is not None:
        QMessageBox.critical(None, "程序发生错误", f"程序没有正常完成操作。\n\n{value}\n\n详细日志：{log_path}")
    else:
        sys.stderr.write(f"程序发生错误：{value}\n详细日志：{log_path}\n")


def self_check() -> int:
    """Non-GUI packaging check used by CI and release validation."""

    marker_text = os.environ.get("CNC_PROGRAM_SHEET_SELF_CHECK_MARKER")
    marker = Path(marker_text) if marker_text else None
    template = resource_path("assets", "CNC程序单.xlsx")
    if marker:
        marker.write_text(f"started\ntemplate={template}\nexists={template.exists()}\n", encoding="utf-8")
    try:
        verify_template(template)
        expected = {"program_name": "B", "tool_type": "C", "tool_number": "I", "diameter": "D", "radius": "E", "machining_data": "F", "depth": "J"}
        if dict(LAYOUT.columns) != expected or LAYOUT.capacity != 19:
            raise TemplateError("内置模板映射校验失败。")
    except Exception as exc:
        if marker:
            marker.write_text(marker.read_text(encoding="utf-8") + f"error={exc}\n", encoding="utf-8")
        raise
    if marker:
        marker.write_text(marker.read_text(encoding="utf-8") + "ok\n", encoding="utf-8")
    return 0


if __name__ == "__main__":
    sys.excepthook = _exception_hook
    raise SystemExit(self_check() if "--self-check" in sys.argv else run())
