from datetime import date
from pathlib import Path

from openpyxl import load_workbook

from cnc_program_sheet.excel_service import (
    LAYOUT,
    automatic_cells,
    automatic_cells_for_rows,
    cleared_program_cells,
    export_workbook,
    output_filename,
    verify_template,
)
from cnc_program_sheet.models import ProgramRecord


def test_official_template_cell_mapping() -> None:
    record = ProgramRecord(
        source_path=Path("O-2.NC"),
        program_name="O-2.NC",
        tool_type="",
        tool_number="T1",
        diameter="D19.5",
        machining_data="G81",
        depth="Z-18.5",
    )
    values = automatic_cells([record], generated_on=date(2026, 10, 7))
    assert values == {"J3": "07 / 10 / 2026", "B7": "O-2.NC", "C7": "", "I7": "T1", "D7": "D19.5", "E7": "", "F7": "G81", "J7": "Z-18.5"}
    assert LAYOUT.capacity == 19


def test_output_never_reuses_existing_name(tmp_path: Path) -> None:
    record = ProgramRecord(Path("O-1.NC"), "O-1.NC")
    first = output_filename([record], tmp_path)
    first.touch()
    second = output_filename([record], tmp_path)
    assert first.name == "CNC程序单_O-1.xlsx"
    assert second.name == "CNC程序单_O-1_2.xlsx"


def test_cleared_program_cells_covers_the_full_import_area() -> None:
    values = cleared_program_cells()
    assert len(values) == LAYOUT.capacity * len(LAYOUT.columns)
    assert values["B7"] == ""
    assert values["I7"] == ""
    assert values["J25"] == ""


def test_automatic_cells_keeps_an_operator_inserted_blank_row() -> None:
    first = ProgramRecord(Path("O-1.NC"), "O-1.NC", tool_number="T1", diameter="D1", machining_data="C40°", depth="Z-0.8")
    second = ProgramRecord(Path("O-2.NC"), "O-2.NC", tool_number="T1", diameter="D19.5", machining_data="G81", depth="Z-18.5")

    values = automatic_cells_for_rows([first, None, second], generated_on=date(2026, 10, 7))

    assert values["B7"] == "O-1.NC"
    assert "B8" not in values
    assert values["B9"] == "O-2.NC"
    assert values["J3"] == "07 / 10 / 2026"


def test_official_template_is_copied_and_only_target_cells_change(tmp_path: Path) -> None:
    template = Path("assets/CNC程序单.xlsx")
    original_bytes = template.read_bytes()
    record = ProgramRecord(
        source_path=Path("ABC123.NC"),
        program_name="ABC123.NC",
        tool_type="",
        tool_number="T7",
        diameter="D19.5",
        machining_data="G83",
        depth="Z-18.5",
    )
    output = tmp_path / "generated.xlsx"
    verify_template(template)
    values = cleared_program_cells()
    values.update(automatic_cells([record], generated_on=date(2026, 10, 7)))
    export_workbook(template, output, values)
    assert template.read_bytes() == original_bytes
    book = load_workbook(output, data_only=False)
    try:
        sheet = book["Sheet1"]
        assert (sheet["B7"].value, sheet["C7"].value, sheet["D7"].value, sheet["F7"].value, sheet["I7"].value, sheet["J7"].value) == (
            "ABC123.NC",
            None,
            "D19.5",
            "G83",
            "T7",
            "Z-18.5",
        )
        assert sheet["A1"].value == "CAM SHEET"
        assert sheet["J3"].value == "07 / 10 / 2026"
        assert sheet["B12"].value is None
        assert sheet["D12"].value is None
        assert sheet["A1"].font.bold
        source_book = load_workbook(template, data_only=False)
        try:
            assert sheet["B7"].style_id == source_book["Sheet1"]["B7"].style_id
        finally:
            source_book.close()
    finally:
        book.close()
