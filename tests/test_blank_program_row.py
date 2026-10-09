import os
from pathlib import Path

os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")

from openpyxl import load_workbook
from PySide6.QtWidgets import QApplication

import cnc_program_sheet.main as app_main
from cnc_program_sheet.excel_service import LAYOUT, export_workbook
from cnc_program_sheet.main import MainWindow
from cnc_program_sheet.models import ProgramRecord


def _record(name: str) -> ProgramRecord:
    return ProgramRecord(
        source_path=Path(name),
        program_name=name,
        tool_number="T1",
        diameter="D1",
        machining_data="C40°",
        depth="Z-0.8",
    )


def test_insert_then_delete_blank_program_row_preserves_export_order(tmp_path: Path) -> None:
    app = QApplication.instance() or QApplication([])
    window = MainWindow()
    template = Path("assets/CNC程序单.xlsx")
    window.load_template(template)
    window.records = [_record("O-1.NC"), _record("O-2.NC")]
    window._apply_records_to_view()

    # Table row 7 is Excel row 8, so this inserts below Excel row 7.
    window.insert_blank_program_row(7)
    assert [record.program_name if record else None for record in window.records] == ["O-1.NC", None, "O-2.NC"]
    assert window.preview.item(6, 1).text() == "O-1.NC"
    assert window.preview.item(7, 1).text() == ""
    assert window.preview.item(8, 1).text() == "O-2.NC"

    output = tmp_path / "spaced.xlsx"
    export_workbook(template, output, window.manual_cells)
    workbook = load_workbook(output, data_only=False)
    try:
        sheet = workbook["Sheet1"]
        assert (sheet["B7"].value, sheet["B8"].value, sheet["B9"].value) == ("O-1.NC", None, "O-2.NC")
    finally:
        workbook.close()

    window.remove_blank_program_row(7)
    assert [record.program_name if record else None for record in window.records] == ["O-1.NC", "O-2.NC"]
    assert window.preview.item(7, 1).text() == "O-2.NC"
    window.close()
    assert app is not None


def test_insert_rotation_marker_is_red_and_exports_as_a_red_instruction(tmp_path: Path) -> None:
    app = QApplication.instance() or QApplication([])
    window = MainWindow()
    template = Path("assets/CNC程序单.xlsx")
    window.load_template(template)
    window.records = [_record("O-1.NC"), _record("O-2.NC")]
    window._apply_records_to_view()

    window.insert_marker_program_row(7, "❮Y⟲180°❯")

    assert [record.program_name if record else None for record in window.records] == ["O-1.NC", None, "O-2.NC"]
    assert window.row_markers == {1: "❮Y⟲180°❯"}
    preview_marker = window.preview.item(7, 1)
    assert preview_marker.text() == "❮Y⟲180°❯"
    assert preview_marker.foreground().color().name() == "#ff0000"

    output = tmp_path / "marker.xlsx"
    export_workbook(template, output, window.manual_cells, red_text_cells=["B8"])
    workbook = load_workbook(output, data_only=False)
    try:
        cell = workbook["Sheet1"]["B8"]
        assert cell.value == "❮Y⟲180°❯"
        assert cell.font.color is not None and cell.font.color.rgb == "FFFF0000"
    finally:
        workbook.close()

    window.remove_blank_program_row(7)
    assert [record.program_name if record else None for record in window.records] == ["O-1.NC", "O-2.NC"]
    assert window.row_markers == {}
    window.close()
    assert app is not None


def test_multi_tool_import_with_available_capacity_does_not_report_negative_overflow(tmp_path: Path, monkeypatch) -> None:
    app = QApplication.instance() or QApplication([])
    source = tmp_path / "0505.NC"
    source.write_text(
        """%
O0000(0505)
(T1|D3R0.2|H1)
(T2|D1|H2)
(T3|D1R0.5|H3)
N100T1M6
N110G1Z-1.
N120T2M6
N130G1Z-2.
N140T3M6
N150G1Z-3.
M30
%""",
        encoding="utf-8",
    )
    window = MainWindow()
    window.load_template(Path("assets/CNC程序单.xlsx"))
    errors: list[tuple[str, str]] = []
    monkeypatch.setattr(window, "show_error", lambda title, details: errors.append((title, details)))

    window.add_nc_files([source])

    assert len(window.records) == 3
    assert errors == []
    window.close()
    assert app is not None


def test_overflow_names_the_exact_program_row_that_did_not_fit(tmp_path: Path, monkeypatch) -> None:
    app = QApplication.instance() or QApplication([])
    source = tmp_path / "too-many.NC"
    source.write_text("O0000\nM30\n", encoding="utf-8")
    parsed = [
        ProgramRecord(
            source_path=source,
            program_name=source.name,
            source_tool_index=index,
            tool_number=f"T{index + 1}",
            diameter="D1",
            machining_data="",
            depth="Z-1",
        )
        for index in range(LAYOUT.capacity + 1)
    ]
    monkeypatch.setattr(app_main, "parse_file_records", lambda _: parsed)
    window = MainWindow()
    window.load_template(Path("assets/CNC程序单.xlsx"))
    reported: list[list[ProgramRecord]] = []
    monkeypatch.setattr(window, "show_capacity_overflow", lambda records: reported.append(records))

    window.add_nc_files([source])

    assert len(window.records) == LAYOUT.capacity
    assert [(record.program_name, record.tool_number) for record in reported[0]] == [("too-many.NC", "T20")]
    window.close()
    assert app is not None
