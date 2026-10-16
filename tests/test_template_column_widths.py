import os
import shutil
from pathlib import Path

os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")

from openpyxl import load_workbook
from PySide6.QtCore import Qt
from PySide6.QtTest import QTest
from PySide6.QtWidgets import QApplication

import cnc_program_sheet.main as main_module
from cnc_program_sheet.excel_service import save_template_column_widths
from cnc_program_sheet.main import ExcelPreview


def _template_copy(tmp_path: Path) -> Path:
    target = tmp_path / "operator-template.xlsx"
    shutil.copy2(Path("assets/CNC程序单.xlsx"), target)
    return target


def test_layout_save_updates_only_column_dimensions_and_reloads_at_same_width(tmp_path: Path) -> None:
    """Ctrl+S layout persistence must not leak preview/program data into a template."""

    template = _template_copy(tmp_path)
    original = load_workbook(template, data_only=False)
    try:
        original_values = (original["Sheet1"]["B7"].value, original["Sheet1"]["J3"].value)
    finally:
        original.close()
    save_template_column_widths(template, {"A": 74, "B": 231, "K": 146})

    workbook = load_workbook(template, data_only=False)
    try:
        sheet = workbook["Sheet1"]
        assert abs(sheet.column_dimensions["B"].width - ((231 - 10 + 0.01) / 7)) < 0.000001
        # The original source template remains a template, not a saved job.
        assert (sheet["B7"].value, sheet["J3"].value) == original_values
    finally:
        workbook.close()

    app = QApplication.instance() or QApplication([])
    preview = ExcelPreview()
    preview.load_template(template)
    try:
        assert preview.columnWidth(0) == 74
        assert preview.columnWidth(1) == 231
        assert preview.columnWidth(10) == 146
        assert not preview.column_widths_dirty
    finally:
        preview.close()
    assert app is not None


def test_ctrl_s_saves_a_dragged_column_width_to_the_current_template(tmp_path: Path, monkeypatch) -> None:
    app = QApplication.instance() or QApplication([])
    template = _template_copy(tmp_path)
    saved_settings: list[dict] = []
    monkeypatch.setattr(main_module, "load_settings", lambda: {})
    monkeypatch.setattr(main_module, "save_settings", lambda settings: saved_settings.append(dict(settings)))
    monkeypatch.setattr(main_module, "default_user_template", lambda: template)

    window = main_module.MainWindow()
    try:
        original_width = window.preview.columnWidth(1)
        expected_width = original_width + 57
        window.preview.setColumnWidth(1, expected_width)
        assert window.preview.column_widths_dirty

        QTest.keyClick(window.preview, Qt.Key.Key_S, Qt.KeyboardModifier.ControlModifier)
        app.processEvents()

        assert not window.preview.column_widths_dirty
        reloaded = ExcelPreview()
        try:
            reloaded.load_template(template)
            assert reloaded.columnWidth(1) == expected_width
        finally:
            reloaded.close()
    finally:
        window.tray_icon.hide()
        window.close()
    assert saved_settings
